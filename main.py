from __future__ import annotations

import csv
import io
import json
import math
import os
import re
import sqlite3
import unicodedata
import uuid
import zipfile
import xml.etree.ElementTree as ET
from contextlib import closing
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

ROOT = Path(__file__).resolve().parent
STATIC_DIR = ROOT / "static"
LOCAL_ZONE = ZoneInfo("Europe/Madrid")
MAX_IMPORT_BYTES = 10 * 1024 * 1024

# Municipio operativo por defecto (solo Castellón por ahora).
DEFAULT_MUNICIPALITY = "Castelló de la Plana"
DEFAULT_PROVINCE = "Castellón"

# Calendario oficial: nacional, Comunitat Valenciana y locales de Castelló de la Plana.
# Los años se calculan por reglas offline (Pascua incluida) desde 2025 hasta cinco
# ejercicios vista. Los festivos nacionales/autonómicos se replican para todos los
# municipios dados de alta; los locales solo para el municipio por defecto.
HOLIDAY_FIRST_YEAR = 2025
HOLIDAY_GENERATED_YEARS_AHEAD = 5


def easter_sunday(year: int) -> date:
    """Domingo de Pascua (algoritmo gregoriano anónimo)."""
    a = year % 19
    b, c = divmod(year, 100)
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = divmod(c, 4)
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    month, day = divmod(h + l - 7 * m + 114, 31)
    return date(year, month, day + 1)


def holidays_for_year(year: int) -> list[tuple[str, str, str]]:
    """Devuelve (fecha, nombre, ámbito) para un ejercicio."""
    easter = easter_sunday(year)
    entries: list[tuple[date, str, str]] = [
        (date(year, 1, 1), "Año Nuevo", "national"),
        (date(year, 1, 6), "Epifanía del Señor", "national"),
        (easter - timedelta(days=2), "Viernes Santo", "national"),
        (date(year, 5, 1), "Fiesta del Trabajo", "national"),
        (date(year, 8, 15), "Asunción de la Virgen", "national"),
        (date(year, 10, 12), "Fiesta Nacional de España", "national"),
        (date(year, 11, 1), "Todos los Santos", "national"),
        (date(year, 12, 6), "Día de la Constitución Española", "national"),
        (date(year, 12, 8), "Inmaculada Concepción", "national"),
        (date(year, 12, 25), "Natividad del Señor", "national"),
        (date(year, 3, 19), "San José", "autonomic"),
        (easter + timedelta(days=1), "Lunes de Pascua", "autonomic"),
        (date(year, 6, 24), "San Juan", "autonomic"),
        (date(year, 10, 9), "Día de la Comunitat Valenciana", "autonomic"),
        # Magdalena: lunes de la tercera semana de Cuaresma (Pascua - 27 días).
        (easter - timedelta(days=27), "Fiestas de la Magdalena (local)", "local"),
    ]
    san_pedro = date(year, 6, 29)
    if san_pedro.weekday() == 6:  # Si cae en domingo, se traslada al lunes.
        san_pedro += timedelta(days=1)
    entries.append((san_pedro, "San Pedro (local)", "local"))
    return [(day.isoformat(), name, scope) for day, name, scope in entries]


def generated_official_holidays(start_year: int, end_year: int) -> list[tuple[str, str, str]]:
    rows: list[tuple[str, str, str]] = []
    for year in range(start_year, end_year + 1):
        rows.extend(holidays_for_year(year))
    return rows


SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS employees (
    id TEXT PRIMARY KEY,
    code TEXT NOT NULL UNIQUE,
    first_name TEXT NOT NULL,
    last_name TEXT NOT NULL DEFAULT '',
    active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS posts (
    id TEXT PRIMARY KEY,
    code TEXT NOT NULL UNIQUE,
    name TEXT NOT NULL,
    municipality TEXT NOT NULL,
    province TEXT NOT NULL,
    active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL,
    UNIQUE(name, municipality)
);
CREATE TABLE IF NOT EXISTS shifts (
    id TEXT PRIMARY KEY,
    employee_id TEXT NOT NULL REFERENCES employees(id),
    post_id TEXT NOT NULL REFERENCES posts(id),
    start_at TEXT NOT NULL,
    end_at TEXT NOT NULL,
    break_start_at TEXT,
    break_end_at TEXT,
    break_paid INTEGER NOT NULL DEFAULT 0 CHECK(break_paid IN (0,1)),
    second_start_at TEXT,
    second_end_at TEXT,
    status TEXT NOT NULL DEFAULT 'confirmed' CHECK(status IN ('pending','confirmed')),
    source TEXT NOT NULL DEFAULT 'manual',
    source_label TEXT,
    created_at TEXT NOT NULL,
    CHECK(end_at > start_at),
    CHECK((break_start_at IS NULL AND break_end_at IS NULL) OR
          (break_start_at IS NOT NULL AND break_end_at IS NOT NULL AND
           break_end_at > break_start_at AND break_start_at >= start_at AND break_end_at <= end_at)),
    CHECK((second_start_at IS NULL AND second_end_at IS NULL) OR
          (second_start_at IS NOT NULL AND second_end_at IS NOT NULL AND
           second_end_at > second_start_at))
);
CREATE INDEX IF NOT EXISTS idx_shifts_start ON shifts(start_at);
CREATE INDEX IF NOT EXISTS idx_shifts_employee_start ON shifts(employee_id, start_at);
CREATE INDEX IF NOT EXISTS idx_shifts_post_start ON shifts(post_id, start_at);
CREATE INDEX IF NOT EXISTS idx_shifts_status ON shifts(status);
CREATE TABLE IF NOT EXISTS holidays (
    id TEXT PRIMARY KEY,
    date TEXT NOT NULL,
    municipality TEXT NOT NULL,
    name TEXT NOT NULL,
    created_at TEXT NOT NULL,
    UNIQUE(date, municipality, name)
);
CREATE TABLE IF NOT EXISTS dismissed_holidays (
    date TEXT NOT NULL,
    municipality TEXT NOT NULL,
    name TEXT NOT NULL,
    PRIMARY KEY(date, municipality, name)
);
CREATE TABLE IF NOT EXISTS review_items (
    id TEXT PRIMARY KEY,
    source_label TEXT NOT NULL,
    raw_json TEXT NOT NULL,
    errors_json TEXT NOT NULL,
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS monthly_summaries (
    id TEXT PRIMARY KEY,
    employee_id TEXT NOT NULL REFERENCES employees(id),
    post_id TEXT REFERENCES posts(id),
    month TEXT NOT NULL,
    minutes INTEGER NOT NULL CHECK(minutes >= 0),
    source TEXT NOT NULL DEFAULT 'manual',
    source_label TEXT,
    created_at TEXT NOT NULL,
    UNIQUE(employee_id, month)
);
CREATE INDEX IF NOT EXISTS idx_monthly_month ON monthly_summaries(month);
"""


def now_text() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def post_code_stem(name: str) -> str:
    text = "".join(char for char in unicodedata.normalize("NFKD", name or "") if not unicodedata.combining(char))
    text = re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-").upper()
    return text[:18] or "CENTRO"


def unique_post_code(used: set[str], name: str, fallback: str = "") -> str:
    base = post_code_stem(name) if name else (fallback or "CENTRO")
    candidate, suffix = base, 2
    while candidate in used:
        candidate = f"{base}-{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def generated_post_code(conn: sqlite3.Connection, name: str) -> str:
    used = {str(row["code"]).upper() for row in conn.execute("SELECT code FROM posts WHERE code IS NOT NULL AND TRIM(code)<>''").fetchall()}
    return unique_post_code(used, name)


def assign_missing_post_codes(conn: sqlite3.Connection) -> None:
    used = {str(row["code"]).upper() for row in conn.execute("SELECT code FROM posts WHERE code IS NOT NULL AND TRIM(code)<>''").fetchall()}
    for row in conn.execute("SELECT id,name FROM posts WHERE code IS NULL OR TRIM(code)=''").fetchall():
        code = unique_post_code(used, row["name"], fallback=f"C-{row['id'][:6].upper()}")
        conn.execute("UPDATE posts SET code=? WHERE id=?", (code, row["id"]))


def connect(db_path: Path) -> sqlite3.Connection:
    conn = sqlite3.connect(db_path, timeout=15)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA busy_timeout = 15000")
    return conn


def sync_official_holidays(conn: sqlite3.Connection) -> int:
    """Carga idempotente de festivos y respeta los que el usuario haya descartado.

    Los nacionales/autonómicos se replican a todos los municipios con centros dados de
    alta; los locales solo se aplican al municipio por defecto (Castelló de la Plana).
    """
    last_year = datetime.now(LOCAL_ZONE).year + HOLIDAY_GENERATED_YEARS_AHEAD
    official = generated_official_holidays(HOLIDAY_FIRST_YEAR, max(HOLIDAY_FIRST_YEAR, last_year))
    dismissed = {(row["date"], row["municipality"], row["name"]) for row in conn.execute("SELECT date,municipality,name FROM dismissed_holidays").fetchall()}
    municipalities = {DEFAULT_MUNICIPALITY}
    municipalities.update(row["municipality"] for row in conn.execute("SELECT DISTINCT municipality FROM posts").fetchall() if row["municipality"])
    rows: list[tuple[str, str, str, str, str]] = []
    for day, name, scope in official:
        targets = municipalities if scope in {"national", "autonomic"} else {DEFAULT_MUNICIPALITY}
        for municipality in targets:
            if (day, municipality, name) in dismissed:
                continue
            rows.append((str(uuid.uuid4()), day, municipality, name, now_text()))
    before = conn.execute("SELECT COUNT(*) FROM holidays").fetchone()[0]
    conn.executemany("INSERT OR IGNORE INTO holidays(id,date,municipality,name,created_at) VALUES(?,?,?,?,?)", rows)
    after = conn.execute("SELECT COUNT(*) FROM holidays").fetchone()[0]
    return after - before


def initialize(db_path: Path) -> None:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    with closing(connect(db_path)) as conn:
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("PRAGMA synchronous=NORMAL")
        conn.executescript(SCHEMA)
        shift_columns = {row["name"] for row in conn.execute("PRAGMA table_info(shifts)").fetchall()}
        if "break_paid" not in shift_columns:
            conn.execute("ALTER TABLE shifts ADD COLUMN break_paid INTEGER NOT NULL DEFAULT 0 CHECK(break_paid IN (0,1))")
        if "second_start_at" not in shift_columns:
            conn.execute("ALTER TABLE shifts ADD COLUMN second_start_at TEXT")
        if "second_end_at" not in shift_columns:
            conn.execute("ALTER TABLE shifts ADD COLUMN second_end_at TEXT")
        post_columns = {row["name"] for row in conn.execute("PRAGMA table_info(posts)").fetchall()}
        if "code" not in post_columns:
            conn.execute("ALTER TABLE posts ADD COLUMN code TEXT")
        assign_missing_post_codes(conn)
        conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_posts_code ON posts(code)")
        sync_official_holidays(conn)
        conn.commit()


class EmployeeInput(BaseModel):
    code: str = Field(min_length=1, max_length=60)
    first_name: str = Field(min_length=1, max_length=100)
    last_name: str = Field(default="", max_length=120)


class PostInput(BaseModel):
    code: str = Field(default="", max_length=40)
    name: str = Field(min_length=1, max_length=160)
    municipality: str = Field(default="", max_length=100)
    province: str = Field(default="", max_length=80)


class ShiftInput(BaseModel):
    employee_id: str
    post_id: str
    start_at: str
    end_at: str
    second_start_at: str | None = None
    second_end_at: str | None = None
    break_start_at: str | None = None
    break_end_at: str | None = None
    break_paid: bool | None = None
    repeat_mode: str = "none"
    repeat_until: str | None = None
    repeat_dates: list[str] | None = None


class ImportRows(BaseModel):
    rows: list[dict[str, Any]] = Field(min_length=1, max_length=5000)


class ShiftIds(BaseModel):
    ids: list[str] = Field(min_length=1, max_length=5000)

class SummaryReviewInput(BaseModel):
    employee_id: str
    month: str
    hours: str


class ConfirmImport(BaseModel):
    rows: list[dict[str, Any]] = Field(min_length=1, max_length=5000)
    filename: str = Field(default="importación", max_length=240)
    source: str = Field(default="csv", max_length=10)


class HolidayInput(BaseModel):
    date: str
    municipality: str = Field(min_length=1, max_length=100)
    name: str = Field(min_length=1, max_length=160)


class MergeInput(BaseModel):
    source_id: str


def parse_local(value: str, label: str) -> datetime:
    try:
        parsed = datetime.fromisoformat(value)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{label}: usa fecha y hora válidas.") from exc
    if parsed.tzinfo is not None:
        parsed = parsed.replace(tzinfo=None)
    candidates = [parsed.replace(tzinfo=LOCAL_ZONE, fold=fold) for fold in (0, 1)]
    round_trips = [item.astimezone(timezone.utc).astimezone(LOCAL_ZONE).replace(tzinfo=None) for item in candidates]
    if any(roundtrip != parsed for roundtrip in round_trips):
        raise ValueError(f"{label}: esa hora local no existe por el cambio horario.")
    if candidates[0].utcoffset() != candidates[1].utcoffset():
        raise ValueError(f"{label}: esa hora local es ambigua por el cambio horario.")
    if parsed.second or parsed.microsecond:
        raise ValueError(f"{label}: la precisión admitida es de un minuto.")
    return parsed.replace(tzinfo=LOCAL_ZONE)


def validate_shift(data: ShiftInput, conn: sqlite3.Connection) -> tuple[datetime, datetime, datetime | None, datetime | None, datetime | None, datetime | None]:
    if not conn.execute("SELECT 1 FROM employees WHERE id=? AND active=1", (data.employee_id,)).fetchone():
        raise ValueError("Selecciona un empleado activo.")
    if not conn.execute("SELECT 1 FROM posts WHERE id=? AND active=1", (data.post_id,)).fetchone():
        raise ValueError("Selecciona un centro activo.")
    start = parse_local(data.start_at, "Entrada")
    end = parse_local(data.end_at, "Salida")
    if end <= start:
        raise ValueError("La salida debe ser posterior a la entrada. Para turnos nocturnos, indica la fecha del día siguiente.")
    if end - start > timedelta(hours=24):
        raise ValueError("Un turno no puede superar 24 horas.")
    if bool(data.second_start_at) != bool(data.second_end_at):
        raise ValueError("Indica entrada y salida del segundo tramo, o deja ambos vacíos.")
    second_start = parse_local(data.second_start_at, "Segunda entrada") if data.second_start_at else None
    second_end = parse_local(data.second_end_at, "Segunda salida") if data.second_end_at else None
    if second_start and second_end:
        if second_start < end:
            raise ValueError("El segundo tramo debe empezar después de la salida del primero.")
        if second_end <= second_start:
            raise ValueError("La segunda salida debe ser posterior a la segunda entrada.")
        if second_end - second_start > timedelta(hours=24):
            raise ValueError("El segundo tramo no puede superar 24 horas.")
        if second_end - start > timedelta(hours=24):
            raise ValueError("La suma del turno partido no puede superar 24 horas.")
    if bool(data.break_start_at) != bool(data.break_end_at):
        raise ValueError("Indica inicio y fin de la pausa, o deja ambos vacíos.")
    break_start = parse_local(data.break_start_at, "Inicio de pausa") if data.break_start_at else None
    break_end = parse_local(data.break_end_at, "Fin de pausa") if data.break_end_at else None
    if break_start and break_end and not (start <= break_start < break_end <= end):
        raise ValueError("La pausa debe estar dentro del turno y terminar después de empezar.")
    if break_start and data.break_paid is None:
        raise ValueError("Confirma si la pausa es remunerada para calcular el turno.")
    return start, end, break_start, break_end, second_start, second_end


def shift_metrics(shift: sqlite3.Row, holiday_dates: set[str], window_start: datetime | None = None, window_end: datetime | None = None) -> dict[str, Any]:
    def field(name: str) -> Any:
        try:
            return shift[name]
        except (KeyError, IndexError):
            return None

    start = parse_local(shift["start_at"], "Entrada").astimezone(timezone.utc)
    end = parse_local(shift["end_at"], "Salida").astimezone(timezone.utc)
    intervals = [(start, end)]
    second_start_raw = field("second_start_at")
    second_end_raw = field("second_end_at")
    if second_start_raw and second_end_raw:
        intervals.append((parse_local(second_start_raw, "Segunda entrada").astimezone(timezone.utc), parse_local(second_end_raw, "Segunda salida").astimezone(timezone.utc)))
    clipped: list[tuple[datetime, datetime]] = []
    for left, right in intervals:
        if window_start:
            left = max(left, window_start.astimezone(timezone.utc))
        if window_end:
            right = min(right, window_end.astimezone(timezone.utc))
        if right > left:
            clipped.append((left, right))
    if not clipped:
        return {"minutes": 0, "night_minutes": 0, "holiday_minutes": 0, "weekend_minutes": 0, "days": {}}
    break_start = parse_local(field("break_start_at"), "Inicio de pausa").astimezone(timezone.utc) if field("break_start_at") and not field("break_paid") else None
    break_end = parse_local(field("break_end_at"), "Fin de pausa").astimezone(timezone.utc) if field("break_end_at") and not field("break_paid") else None
    totals = {"minutes": 0, "night_minutes": 0, "holiday_minutes": 0, "weekend_minutes": 0}
    days: dict[str, dict[str, int]] = {}
    for start, end in clipped:
        boundaries = {start, end}
        first_day = start.astimezone(LOCAL_ZONE).date() - timedelta(days=1)
        final_day = end.astimezone(LOCAL_ZONE).date() + timedelta(days=1)
        current_day = first_day
        while current_day <= final_day:
            for boundary_time in (time.min, time(6), time(22)):
                boundary = datetime.combine(current_day, boundary_time, LOCAL_ZONE).astimezone(timezone.utc)
                if start < boundary < end:
                    boundaries.add(boundary)
            current_day += timedelta(days=1)
        if break_start and start < break_start < end:
            boundaries.add(break_start)
        if break_end and start < break_end < end:
            boundaries.add(break_end)
        points = sorted(boundaries)
        for left, right in zip(points, points[1:]):
            if break_start and break_start <= left < break_end:
                continue
            minute_count = int((right - left).total_seconds() // 60)
            if minute_count <= 0:
                continue
            local = left.astimezone(LOCAL_ZONE)
            day = local.date().isoformat()
            is_night = local.hour >= 22 or local.hour < 6
            is_holiday = day in holiday_dates
            is_weekend = local.weekday() >= 5
            totals["minutes"] += minute_count
            if is_night:
                totals["night_minutes"] += minute_count
            if is_holiday:
                totals["holiday_minutes"] += minute_count
            if is_weekend:
                totals["weekend_minutes"] += minute_count
            bucket = days.setdefault(day, {"minutes": 0, "night_minutes": 0, "holiday_minutes": 0})
            bucket["minutes"] += minute_count
            bucket["night_minutes"] += minute_count if is_night else 0
            bucket["holiday_minutes"] += minute_count if is_holiday else 0
    totals["days"] = days
    return totals


def hhmm(minutes: int) -> str:
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def safe_export_text(value: Any) -> str:
    """Evita que un texto importado se evalúe como fórmula al abrir CSV/Excel."""
    text = str(value or "")
    return f"'{text}" if text[:1] in {"=", "+", "-", "@"} else text


def normalized_name(*parts: str) -> str:
    """Nombre comparable: minúsculas, sin acentos ni espacios sobrantes."""
    text = " ".join(part for part in parts if part).strip().lower()
    return "".join(char for char in unicodedata.normalize("NFKD", text) if not unicodedata.combining(char))


def row_dict(row: sqlite3.Row) -> dict[str, Any]:
    return dict(row)


def cell_to_text(value: Any) -> str:
    """Normaliza celdas de CSV/Excel/PDF a texto estable para el parser de alias."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, timedelta):
        total_minutes = int(value.total_seconds() // 60)
        return f"{total_minutes // 60}:{total_minutes % 60:02d}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def coerce_date(value: str) -> date:
    text = (value or "").strip()
    if not text:
        raise ValueError("falta la fecha")
    try:
        return date.fromisoformat(text)
    except ValueError:
        pass
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        pass
    if re.fullmatch(r"\d{8}", text):
        try:
            return date(int(text[:4]), int(text[4:6]), int(text[6:8]))
        except ValueError:
            pass
    try:
        serial = float(text.replace(",", "."))
        if 20000 <= serial <= 80000:  # Fecha serial de Excel (base 1899-12-30).
            return date(1899, 12, 30) + timedelta(days=int(serial))
    except ValueError:
        pass
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d", "%d/%m/%y", "%d-%m-%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    month_match = re.fullmatch(r"(\d{1,2})[\s/.-]+([a-zA-ZáéíóúÁÉÍÓÚüÜ]+)[\s/.-]+(\d{4})", text)
    if month_match:
        month_key = "".join(char for char in unicodedata.normalize("NFKD", month_match.group(2).lower()) if not unicodedata.combining(char))
        month = MONTH_NAMES.get(month_key)
        if month:
            try:
                return date(int(month_match.group(3)), month, int(month_match.group(1)))
            except ValueError:
                pass
    raise ValueError("fecha no válida (usa YYYY-MM-DD o DD/MM/YYYY)")


def coerce_time(value: str) -> time:
    text = (value or "").strip().replace("\u00a0", " ")
    if not text:
        raise ValueError("falta la hora")
    try:
        return time.fromisoformat(text)
    except ValueError:
        pass
    try:
        return datetime.fromisoformat(text).time()
    except ValueError:
        pass
    match = re.fullmatch(r"(\d{1,2})[:.](\d{2})(?::(\d{2}))?\s*(?:h|hs)?", text, re.IGNORECASE)
    if match:
        hour, minute = int(match.group(1)), int(match.group(2))
        second = int(match.group(3) or 0)
        if hour == 24 and minute == 0 and second == 0:
            hour = 0
        if 0 <= hour < 24 and 0 <= minute < 60 and 0 <= second < 60:
            return time(hour, minute, second)
    try:
        numeric = float(text.replace(",", "."))
        if 0 <= numeric < 1:  # Fracción de día típica de Excel (0.5 = 12:00).
            total = round(numeric * 1440)
            return time((total // 60) % 24, total % 60)
        if numeric.is_integer() and 0 <= numeric < 24:
            return time(int(numeric), 0)
    except ValueError:
        pass
    raise ValueError("hora no válida (usa HH:MM)")


def import_time(raw_date: str, raw_time: str, base_date: date, start: datetime | None = None) -> datetime:
    text = (raw_time or "").strip()
    if re.search(r"\d{4}-\d{2}-\d{2}", text) and re.search(r"\d{1,2}:\d{2}", text):
        try:
            parsed = datetime.fromisoformat(text.replace(" ", "T"))
            return parsed.replace(tzinfo=None)
        except ValueError:
            pass
    embedded = re.fullmatch(r"(.+?)[ T](\d{1,2}[:.]\d{2}(?::\d{2})?\s*(?:h|hs)?)", text, re.IGNORECASE)
    if embedded:
        try:
            return datetime.combine(coerce_date(embedded.group(1)), coerce_time(embedded.group(2)))
        except ValueError:
            pass
    parsed_time = coerce_time(text)
    parsed_date = coerce_date(raw_date) if (raw_date or "").strip() else base_date
    result = datetime.combine(parsed_date, parsed_time)
    if start and result <= start:
        result += timedelta(days=1)
    return result


IMPORT_ALIASES = {
    "id": "code", "employee_id": "code", "employee_code": "code", "codigo": "code", "código": "code",
    "codigo empleado": "code", "código empleado": "code", "codigo trabajador": "code", "código trabajador": "code", "cod empleado": "code", "cod. empleado": "code",
    "id trabajador": "code", "id empleado": "code", "n empleado": "code", "nº empleado": "code",
    "n.º empleado": "code", "num empleado": "code", "numero empleado": "code", "número empleado": "code",
    "employee": "employee", "empleado": "employee", "nombre": "employee", "nombre ficticio": "employee",
    "nombre completo": "employee", "nombre y apellidos": "employee", "trabajador": "employee", "vigilante": "employee",
    "apellidos": "last_name", "apellido": "last_name",
    "post": "post", "service": "post", "puesto": "post", "puesto de trabajo": "post", "servicio": "post",
    "instalacion": "post", "instalación": "post", "servicio/instalacion": "post", "servicio/instalación": "post",
    "centro": "post", "centro de trabajo": "post",
    "post_code": "post_code", "id centro": "post_code", "id del centro": "post_code", "id de centro": "post_code",
    "codigo centro": "post_code", "código centro": "post_code", "cod centro": "post_code", "cod. centro": "post_code",
    "municipality": "municipality", "municipio": "municipality", "localidad": "municipality", "poblacion": "municipality",
    "población": "municipality", "ciudad": "municipality",
    "province": "province", "provincia": "province", "prov": "province",
    "date": "date", "fecha": "date", "dia": "date", "día": "date", "fecha turno": "date",
    "fecha de trabajo": "date", "fecha servicio": "date",
    "start": "start", "entrada": "start", "hora entrada": "start", "hora de entrada": "start", "inicio": "start",
    "hora inicio": "start", "fecha entrada": "start", "fecha de entrada": "start", "fecha inicio": "start",
    "fecha/hora entrada": "start",
    "end": "end", "salida": "end", "hora salida": "end", "hora de salida": "end", "fin": "end",
    "hora fin": "end", "fecha salida": "end", "fecha de salida": "end", "fecha fin": "end", "fecha/hora salida": "end",
    "schedule": "schedule", "horario": "schedule", "horario de trabajo": "schedule", "franja horaria": "schedule",
    "franja": "schedule", "entrada salida": "schedule", "entrada/salida": "schedule", "horario turno": "schedule",
    "second_start": "second_start", "segunda entrada": "second_start", "entrada 2": "second_start",
    "segundo tramo inicio": "second_start", "inicio segundo tramo": "second_start", "tramo 2 inicio": "second_start",
    "second_end": "second_end", "segunda salida": "second_end", "salida 2": "second_end",
    "segundo tramo fin": "second_end", "fin segundo tramo": "second_end", "tramo 2 fin": "second_end",
    "break_start": "break_start", "pausa inicio": "break_start", "inicio pausa": "break_start",
    "inicio descanso": "break_start", "descanso inicio": "break_start", "comida inicio": "break_start",
    "break_end": "break_end", "pausa fin": "break_end", "fin pausa": "break_end",
    "fin descanso": "break_end", "descanso fin": "break_end", "comida fin": "break_end",
    "free_start": "free_start", "espacio libre desde": "free_start", "inicio espacio libre": "free_start",
    "libre desde": "free_start", "salida espacio libre": "free_start",
    "free_end": "free_end", "espacio libre hasta": "free_end", "fin espacio libre": "free_end",
    "libre hasta": "free_end", "entrada espacio libre": "free_end",
    "break_paid": "break_paid", "pausa remunerada": "break_paid", "pausa pagada": "break_paid",
    "location": "location", "localizacion": "location", "localización": "location",
    "localización ficticia": "location", "localizacion ficticia": "location",
    "month": "month", "mes": "month", "periodo": "month", "periodo/mes": "month", "mes/año": "month",
    "total_hours": "total_hours", "horas": "total_hours", "horas totales": "total_hours",
    "horas totales trabajadas": "total_hours", "horas totales mes": "total_hours", "horas mes": "total_hours",
}

MONTH_NAMES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
}


def normalized_header(value: Any) -> str:
    text = cell_to_text(value).lower()
    text = "".join(char for char in unicodedata.normalize("NFKD", text) if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


HEADER_ALIASES = {normalized_header(alias): canonical for alias, canonical in IMPORT_ALIASES.items()}


def canonical_header(value: Any) -> str | None:
    """Resuelve encabezados reales (con acentos, símbolos y variantes) a un campo interno."""
    key = normalized_header(value)
    if not key:
        return None
    if key in HEADER_ALIASES:
        return HEADER_ALIASES[key]
    padded = f" {key} "
    for alias in sorted(HEADER_ALIASES, key=len, reverse=True):
        if len(alias) >= 3 and f" {alias} " in padded:
            return HEADER_ALIASES[alias]
    tokens = set(key.split())
    if ("centro" in tokens or "puesto" in tokens or "servicio" in tokens or "instalacion" in tokens):
        if tokens & {"id", "codigo", "cod", "referencia", "ref"}:
            return "post_code"
        return "post"
    if tokens & {"dni", "nif", "nie", "documento"}:
        return "code"
    if tokens & {"id", "codigo", "cod", "referencia", "ref"} and tokens & {"empleado", "trabajador", "vigilante", "persona"}:
        return "code"
    if tokens & {"nombre", "apellidos"} and tokens & {"empleado", "trabajador", "vigilante", "persona"}:
        return "employee"
    if tokens & {"empleado", "trabajador", "vigilante", "persona"}:
        return "employee"
    if ("total" in tokens or "totales" in tokens) and ("hora" in tokens or "horas" in tokens):
        return "total_hours"
    if ("mes" in tokens or "ano" in tokens) and ("hora" in tokens or "horas" in tokens):
        return "total_hours"
    if key in {"horas", "total", "total horas", "horas totales", "total mes", "horas mes"}:
        return "total_hours"
    if "fecha" in tokens or tokens & {"dia", "fecha"}:
        return "date"
    if "libre" in tokens:
        if tokens & {"desde", "inicio", "entrada", "salida"}:
            return "free_start"
        if tokens & {"hasta", "fin"}:
            return "free_end"
        return "free_start"
    if "entrada" in tokens or "inicio" in tokens:
        return "start"
    if "salida" in tokens or "fin" in tokens:
        return "end"
    return None


def parse_month(value: str) -> str:
    text = (value or "").strip()
    if not text:
        raise ValueError("falta el mes")
    match = re.fullmatch(r"(\d{4})-(\d{1,2})(?:-\d{1,2})?(?:[ T].*)?", text)
    if match:
        month = int(match.group(2))
        if not 1 <= month <= 12:
            raise ValueError("mes no válido")
        return f"{int(match.group(1)):04d}-{month:02d}"
    match = re.fullmatch(r"(\d{1,2})[/-](\d{4})", text)
    if match:
        month = int(match.group(1))
        if not 1 <= month <= 12:
            raise ValueError("mes no válido")
        return f"{int(match.group(2)):04d}-{month:02d}"
    tokens = re.split(r"[\s/,.]+", text.lower())
    year = next((int(token) for token in tokens if re.fullmatch(r"\d{4}", token)), None)
    name = next((token for token in tokens if token in MONTH_NAMES), None)
    if year and name:
        return f"{year:04d}-{MONTH_NAMES[name]:02d}"
    raise ValueError(f"mes no válido ({value})")


def parse_hours(value: str) -> int:
    text = (value or "").strip().replace(",", ".")
    text = re.sub(r"\s*h(?:oras?)?$", "", text, flags=re.I).strip()
    if not text:
        raise ValueError("faltan las horas")
    try:
        if ":" in text:
            hours, _, minutes = text.partition(":")
            total = int(hours) * 60 + int((minutes or "0")[:2] or 0)
        else:
            number = float(text)
            if not math.isfinite(number):
                raise ValueError("las horas deben ser un número finito")
            total = round(number * 60)
    except (ValueError, OverflowError) as exc:
        raise ValueError(f"horas no válidas ({value})") from exc
    if total < 0 or total > 100000:
        raise ValueError("las horas deben estar entre 0 y 100000 minutos")
    return total


def resolve_or_create_employee(conn: sqlite3.Connection, row: dict[str, Any]) -> str:
    employee_id = str(row.get("employee_id") or "")
    if employee_id and conn.execute("SELECT 1 FROM employees WHERE id=?", (employee_id,)).fetchone():
        return employee_id
    code = str(row.get("employee_code") or "").strip()
    if not code:
        return ""
    existing = conn.execute("SELECT id FROM employees WHERE code=?", (code,)).fetchone()
    if existing:
        return existing["id"]
    new_id = str(uuid.uuid4())
    name = str(row.get("employee_name") or "").strip()
    first, _, last = name.partition(" ")
    conn.execute(
        "INSERT INTO employees(id,code,first_name,last_name,created_at) VALUES(?,?,?,?,?)",
        (new_id, code, first or code, last, now_text()),
    )
    return new_id


def resolve_or_create_post(conn: sqlite3.Connection, row: dict[str, Any]) -> str:
    post_id = str(row.get("post_id") or "")
    if post_id and conn.execute("SELECT 1 FROM posts WHERE id=?", (post_id,)).fetchone():
        return post_id
    code = str(row.get("post_code") or "").strip().upper()
    if code:
        existing = conn.execute("SELECT id FROM posts WHERE UPPER(code)=?", (code,)).fetchone()
        if existing:
            return existing["id"]
    name = str(row.get("post_name") or "").strip()
    municipality = str(row.get("municipality") or "").strip() or DEFAULT_MUNICIPALITY
    if name:
        existing = conn.execute("SELECT id FROM posts WHERE name=? AND municipality=?", (name, municipality)).fetchone()
        if existing:
            return existing["id"]
        candidates = conn.execute("SELECT id FROM posts WHERE name=? ORDER BY active DESC LIMIT 2", (name,)).fetchall()
        if len(candidates) == 1:
            return candidates[0]["id"]
    if not name and not code:
        return ""
    if not name:
        name = code
    province = str(row.get("province") or "").strip() or DEFAULT_PROVINCE
    new_id = str(uuid.uuid4())
    if not code:
        code = generated_post_code(conn, name)
    conn.execute(
        "INSERT INTO posts(id,code,name,municipality,province,created_at) VALUES(?,?,?,?,?,?)",
        (new_id, code, name, municipality, province, now_text()),
    )
    sync_official_holidays(conn)
    return new_id


def queue_review(conn: sqlite3.Connection, label: str, row: dict[str, Any], errors: list[str], now: str) -> None:
    conn.execute(
        "INSERT INTO review_items(id,source_label,raw_json,errors_json,created_at) VALUES(?,?,?,?,?)",
        (str(uuid.uuid4()), label, json.dumps(row, ensure_ascii=False), json.dumps(errors or ["Revisa la fila antes de incluirla en los totales."], ensure_ascii=False), now),
    )


def store_summary_row(conn: sqlite3.Connection, row: dict[str, Any], errors: list[str], label: str, source: str, filename: str, now: str) -> int:
    """Guarda/actualiza un resumen mensual (no son turnos: no entran en los totales de turnos)."""
    if errors:
        queue_review(conn, label, row, errors, now)
        return 0
    employee_id = resolve_or_create_employee(conn, row)
    if not employee_id:
        errors.append("No se pudo identificar o crear el empleado (falta el ID).")
    month = str(row.get("month") or "")
    minutes = row.get("minutes")
    if not month:
        errors.append("Mes no válido.")
    if minutes is None:
        errors.append("Horas no válidas.")
    if errors:
        queue_review(conn, label, row, errors, now)
        return 0
    post_id = resolve_or_create_post(conn, row) or None
    existing = conn.execute("SELECT id FROM monthly_summaries WHERE employee_id=? AND month=?", (employee_id, month)).fetchone()
    if existing:
        conn.execute("UPDATE monthly_summaries SET post_id=?, minutes=?, source=?, source_label=?, created_at=? WHERE id=?", (post_id, int(minutes), source, Path(filename).name, now, existing["id"]))
    else:
        conn.execute("INSERT INTO monthly_summaries(id,employee_id,post_id,month,minutes,source,source_label,created_at) VALUES(?,?,?,?,?,?,?,?)", (str(uuid.uuid4()), employee_id, post_id, month, int(minutes), source, Path(filename).name, now))
    return 1

def build_name_index(conn: sqlite3.Connection) -> dict[str, list[str]]:
    index: dict[str, list[str]] = {}
    for row in conn.execute("SELECT code, first_name, last_name FROM employees WHERE active=1").fetchall():
        key = normalized_name(row["first_name"], row["last_name"])
        if key:
            index.setdefault(key, []).append(row["code"])
    return index


def combined_employee_name(fields: dict[str, str]) -> str:
    parts = [fields.get("employee", "").strip(), fields.get("last_name", "").strip()]
    return " ".join(part for part in parts if part)


def split_code_name(value: str) -> tuple[str, str] | None:
    """Separa celdas compuestas tipo «C-01 - Centro Norte», «VS-001 · Ana» o «Centro (C-01)»."""
    text = (value or "").strip()
    if not text:
        return None
    patterns = [
        r"^\s*([A-Za-z]{0,4}[-.]?\d[A-Za-z0-9._/-]{0,18})\s*[-–—:·|]\s*(.+)$",
        r"^\s*(.+?)\s*\(\s*([A-Za-z]{0,4}[-.]?\d[A-Za-z0-9._/-]{0,18})\s*\)\s*$",
        r"^\s*([A-Za-z]{0,4}[-.]?\d[A-Za-z0-9._/-]{0,18})\s+(.+)$",
        r"^\s*(.+?)\s*[-–—:·|]\s*([A-Za-z]{0,4}[-.]?\d[A-Za-z0-9._/-]{0,18})\s*$",
    ]
    for pattern in patterns:
        match = re.fullmatch(pattern, text)
        if not match:
            continue
        first, second = match.group(1).strip(), match.group(2).strip()
        if re.search(r"\d", first) and not re.search(r"\d", second[:4]):
            return first.upper(), second
        if re.search(r"\d", second) and not re.search(r"\d", first[:4]):
            return second.upper(), first
    return None


def parse_schedule_range(value: str) -> dict[str, str] | None:
    """Interpreta «08:00-16:00», «08:00 a 14:00 y 17:00 a 20:00» o «22:00-06:00»."""
    text = (value or "").strip()
    if not text:
        return None
    colon_tokens = re.findall(r"(?<!\d)(\d{1,2})[:.](\d{2})(?!\d)", text)
    if colon_tokens:
        minutes = [int(hour) * 60 + int(minute) for hour, minute in colon_tokens]
    else:
        plain_tokens = re.findall(r"(?<!\d)(\d{1,2})(?:\s*h(?:s|oras?)?)?(?!\d)", text)
        minutes = [int(hour) * 60 for hour in plain_tokens]
    minutes = [minute for minute in minutes if 0 <= minute < 24 * 60]
    if len(minutes) == 2:
        start, end = minutes
        if start == end:
            return None
        return {"start": f"{start // 60:02d}:{start % 60:02d}", "end": f"{end // 60:02d}:{end % 60:02d}"}
    if len(minutes) == 4:
        start = minutes[0]
        first_end = minutes[1] + (24 * 60 if minutes[1] <= start else 0)
        second_start = minutes[2] + (24 * 60 if minutes[2] <= first_end else 0)
        end = minutes[3] + (24 * 60 if minutes[3] <= second_start else 0)
        if not (start < first_end and second_start < end) or first_end > second_start:
            return None
        if first_end == second_start:
            return {"start": f"{start // 60:02d}:{start % 60:02d}", "end": f"{end % (24 * 60) // 60:02d}:{end % 60:02d}"}
        return {
            "start": f"{start // 60:02d}:{start % 60:02d}",
            "end": f"{end % (24 * 60) // 60:02d}:{end % 60:02d}",
            "free_start": f"{first_end % (24 * 60) // 60:02d}:{first_end % 60:02d}",
            "free_end": f"{second_start % (24 * 60) // 60:02d}:{second_start % 60:02d}",
        }
    return None


def normalize_rows(raw_rows: list[dict[str, Any]], conn: sqlite3.Connection) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    name_index = build_name_index(conn)
    for line_number, raw in enumerate(raw_rows, start=2):
        fields = {(canonical_header(key) or header_key(key)): cell_to_text(value) for key, value in raw.items() if key}
        code_value = (fields.get("code") or "").strip().upper()
        if not any(fields.get(key) for key in ("date", "start", "end")) and fields.get("total_hours") and not fields.get("employee") and (not code_value or code_value in {"TOTAL", "TOTALES", "SUBTOTAL", "SUMA", "SUMAN"}):
            continue
        looks_like_shift = any(fields.get(key) for key in ("date", "start", "end"))
        if (fields.get("month") or fields.get("total_hours")) and not looks_like_shift:
            record = normalize_summary_row(fields, line_number, conn, name_index)
        else:
            record = normalize_shift_row(fields, line_number, conn, name_index)
        code = str(record.get("employee_code") or "").strip()
        name = str(record.get("employee_name") or "").strip()
        if code and name:
            key = normalized_name(name)
            if key and code not in name_index.get(key, []):
                name_index.setdefault(key, []).append(code)
        records.append(record)
    return records


def normalize_shift_row(fields: dict[str, str], line_number: int, conn: sqlite3.Connection, name_index: dict[str, list[str]]) -> dict[str, Any]:
    errors: list[str] = []
    code = fields.get("code", "").strip()
    name = combined_employee_name(fields)
    if code and not name:
        split = split_code_name(code)
        if split:
            code, name = split
    if not code and name:
        split = split_code_name(name)
        if split:
            code, name = split
    if not code and name:
        name_matches = name_index.get(normalized_name(name), [])
        if len(name_matches) == 1:
            code = name_matches[0]
        elif len(name_matches) > 1:
            errors.append(f"fila {line_number}: nombre duplicado ({name}) coincide con {'/'.join(name_matches)}; revisa y unifica el empleado")
    employee = conn.execute("SELECT * FROM employees WHERE code=? AND active=1", (code,)).fetchone() if code else None
    duplicate_of: list[str] = []
    if not employee and name:
        duplicate_of = [candidate for candidate in name_index.get(normalized_name(name), []) if candidate != code]
        if duplicate_of:
            errors.append(f"fila {line_number}: nombre duplicado ({name}) coincide con {'/'.join(duplicate_of)}; revisa y unifica el empleado")
    if not code:
        errors.append(f"fila {line_number}: falta el ID de empleado y el nombre no identifica a un único trabajador")
    post_name = fields.get("post", "").strip()
    post_code = fields.get("post_code", "").strip().upper()
    if post_name and not post_code:
        split = split_code_name(post_name)
        if split:
            post_code, post_name = split
    municipality = fields.get("municipality", "").strip() or DEFAULT_MUNICIPALITY
    province = fields.get("province", "").strip() or DEFAULT_PROVINCE
    post = conn.execute("SELECT * FROM posts WHERE UPPER(code)=? AND active=1", (post_code,)).fetchone() if post_code else None
    if not post and post_name:
        post = conn.execute("SELECT * FROM posts WHERE name=? AND municipality=? AND active=1", (post_name, municipality)).fetchone()
        if not post:
            candidates = conn.execute("SELECT * FROM posts WHERE name=? AND active=1", (post_name,)).fetchall()
            if len(candidates) == 1:
                post = candidates[0]
            elif len(candidates) > 1:
                errors.append(f"fila {line_number}: el nombre de centro «{post_name}» coincide con varios centros; indica su ID")
    if not post_name and not post_code:
        errors.append(f"fila {line_number}: falta el centro")
    if not post_name and post:
        post_name = post["name"]
    date_text = fields.get("date", "")
    start_text = fields.get("start", "")
    end_text = fields.get("end", "")
    schedule_text = fields.get("schedule", "").strip()
    if schedule_text and (not start_text.strip() or not end_text.strip() or not (fields.get("free_start") or fields.get("free_end"))):
        parsed_schedule = parse_schedule_range(schedule_text)
        if parsed_schedule:
            start_text = start_text.strip() or parsed_schedule["start"]
            end_text = end_text.strip() or parsed_schedule["end"]
            if parsed_schedule.get("free_start") and not (fields.get("free_start") or fields.get("free_end")):
                fields["free_start"] = parsed_schedule["free_start"]
                fields["free_end"] = parsed_schedule["free_end"]
        elif not start_text.strip() or not end_text.strip():
            errors.append(f"fila {line_number}: no se pudo interpretar el horario «{schedule_text}»")
    if not date_text.strip():
        try:
            date_text = datetime.fromisoformat(start_text.strip()).date().isoformat()
        except (ValueError, TypeError):
            date_text = ""
    try:
        day = coerce_date(date_text)
        start = import_time(day.isoformat(), start_text, day)
        end = import_time(day.isoformat(), end_text, day, start)
        if end <= start:
            raise ValueError("la salida debe ser posterior a la entrada")
        if end - start > timedelta(hours=24):
            raise ValueError("el turno supera 24 horas")
    except ValueError as exc:
        errors.append(f"fila {line_number}: fecha/horas: {exc}")
        day, start, end = date.today(), datetime.combine(date.today(), time(0)), datetime.combine(date.today(), time(1))
    break_start = break_end = None
    free_mode = bool(fields.get("free_start") or fields.get("free_end"))
    if free_mode and (fields.get("break_start") or fields.get("break_end")):
        errors.append(f"fila {line_number}: usa espacio libre o pausa, no ambos")
    elif free_mode:
        if bool(fields.get("free_start")) != bool(fields.get("free_end")):
            errors.append(f"fila {line_number}: indica inicio y fin del espacio libre")
        elif not any("fecha/horas:" in message for message in errors):
            try:
                break_start = import_time(day.isoformat(), fields.get("free_start", ""), day, start)
                break_end = import_time(day.isoformat(), fields.get("free_end", ""), day, break_start)
                if not (start <= break_start < break_end <= end):
                    raise ValueError("el espacio libre debe estar dentro del turno")
            except ValueError as exc:
                errors.append(f"fila {line_number}: espacio libre: {exc}")
    elif fields.get("break_start") or fields.get("break_end"):
        try:
            break_start = import_time(day.isoformat(), fields.get("break_start", ""), day, start)
            break_end = import_time(day.isoformat(), fields.get("break_end", ""), day, break_start)
            if not (start <= break_start < break_end <= end):
                raise ValueError("la pausa debe estar dentro del turno")
        except ValueError as exc:
            errors.append(f"fila {line_number}: pausa: {exc}")
    paid_raw = fields.get("break_paid", "").lower()
    if free_mode:
        paid_raw = "false"
    elif (break_start or break_end) and paid_raw not in {"si", "sí", "no", "true", "false", "1", "0"}:
        errors.append(f"fila {line_number}: indica si la pausa es remunerada (sí/no)")
    if not free_mode and bool(fields.get("break_start")) != bool(fields.get("break_end")):
        errors.append(f"fila {line_number}: indica inicio y fin de pausa")
    second_start = second_end = None
    if bool(fields.get("second_start")) != bool(fields.get("second_end")):
        errors.append(f"fila {line_number}: indica entrada y salida del segundo tramo")
    elif fields.get("second_start") and fields.get("second_end") and not any("fecha/horas:" in message for message in errors):
        try:
            second_start = import_time(day.isoformat(), fields.get("second_start", ""), day, end)
            second_end = import_time(day.isoformat(), fields.get("second_end", ""), day, second_start)
            if second_start < end or second_end <= second_start:
                raise ValueError("el segundo tramo debe empezar después del primero y terminar después de empezar")
            if second_end - start > timedelta(hours=24):
                raise ValueError("la suma del turno partido supera 24 horas")
        except ValueError as exc:
            second_start = second_end = None
            errors.append(f"fila {line_number}: segundo tramo: {exc}")
    if employee:
        employee_label = f"{employee['code']} · {employee['first_name']} {employee['last_name']}".strip()
    else:
        employee_label = name or code or "Empleado sin ID"
    return {
        "kind": "shift",
        "line": line_number,
        "employee_id": employee["id"] if employee else "",
        "employee_code": code,
        "employee_name": name,
        "employee_label": employee_label,
        "post_id": post["id"] if post else "",
        "post_code": post["code"] if post else post_code,
        "post_name": post_name,
        "municipality": municipality,
        "province": province,
        "post_label": f"{post['code']} · {post_name}" if post else (f"{post_code} · {post_name}" if post_code else (post_name or "Centro sin indicar")),
        "start_at": "" if any("fecha/horas:" in message for message in errors) else start.isoformat(timespec="minutes"),
        "end_at": "" if any("fecha/horas:" in message for message in errors) else end.isoformat(timespec="minutes"),
        "second_start_at": second_start.isoformat(timespec="minutes") if second_start else None,
        "second_end_at": second_end.isoformat(timespec="minutes") if second_end else None,
        "break_start_at": break_start.isoformat(timespec="minutes") if break_start else None,
        "break_end_at": break_end.isoformat(timespec="minutes") if break_end else None,
        "break_paid": paid_raw in {"si", "sí", "true", "1"} if paid_raw else (False if not break_start else None),
        "duplicate_of": duplicate_of,
        "errors": errors,
        "status": "pending" if errors else "ready",
    }


def normalize_summary_row(fields: dict[str, str], line_number: int, conn: sqlite3.Connection, name_index: dict[str, list[str]]) -> dict[str, Any]:
    errors: list[str] = []
    code = fields.get("code", "").strip()
    name = combined_employee_name(fields)
    if code and not name:
        split = split_code_name(code)
        if split:
            code, name = split
    if not code and name:
        split = split_code_name(name)
        if split:
            code, name = split
    employee = conn.execute("SELECT * FROM employees WHERE code=? AND active=1", (code,)).fetchone() if code else None
    duplicate_of: list[str] = []
    if not code and name:
        name_matches = name_index.get(normalized_name(name), [])
        if len(name_matches) == 1:
            code = name_matches[0]
            employee = conn.execute("SELECT * FROM employees WHERE code=? AND active=1", (code,)).fetchone()
        elif len(name_matches) > 1:
            duplicate_of = name_matches
            errors.append(f"fila {line_number}: nombre duplicado ({name}) coincide con {'/'.join(name_matches)}; revisa y unifica el empleado")
    if not employee and name:
        duplicate_of = [candidate for candidate in name_index.get(normalized_name(name), []) if candidate != code]
        if duplicate_of:
            errors.append(f"fila {line_number}: nombre duplicado ({name}) coincide con {'/'.join(duplicate_of)}; revisa y unifica el empleado")
    if not code:
        errors.append(f"fila {line_number}: falta el ID de empleado y el nombre no identifica a un único trabajador")
    try:
        month = parse_month(fields.get("month", ""))
    except ValueError as exc:
        errors.append(f"fila {line_number}: mes: {exc}")
        month = ""
    try:
        minutes = parse_hours(fields.get("total_hours", ""))
    except ValueError as exc:
        errors.append(f"fila {line_number}: horas: {exc}")
        minutes = None
    post_name = (fields.get("location") or fields.get("post") or "").strip()
    return {
        "kind": "summary",
        "line": line_number,
        "employee_id": employee["id"] if employee else "",
        "employee_code": code,
        "employee_name": name,
        "employee_label": name or code or "Empleado sin ID",
        "post_name": post_name,
        "municipality": fields.get("municipality", "").strip() or DEFAULT_MUNICIPALITY,
        "province": fields.get("province", "").strip() or DEFAULT_PROVINCE,
        "month": month,
        "minutes": minutes,
        "hours": hhmm(minutes) if minutes is not None else "",
        "duplicate_of": duplicate_of,
        "errors": errors,
        "status": "pending" if errors else "ready",
    }


def header_key(value: Any) -> str:
    return " ".join(cell_to_text(value).lower().split())


def header_score(row: list[Any]) -> int:
    return sum(1 for cell in row if canonical_header(cell))


def matrix_header_index(rows: list[list[Any]]) -> int:
    """Busca la fila de encabezados aunque haya títulos o filas vacías delante."""
    best_score, best_index = 0, None
    for index, row in enumerate(rows[:20]):
        score = header_score(row)
        if score > best_score:
            best_score, best_index = score, index
    if best_index is not None and best_score >= 2:
        return best_index
    for index, row in enumerate(rows):
        if any(cell_to_text(cell) for cell in row):
            return index
    return 0


def matrix_to_records(matrix: list[list[Any]]) -> list[dict[str, str]]:
    """Convierte una matriz (CSV/Excel/PDF/ODS) en filas usando la fila de encabezados detectada."""
    rows = [list(row) for row in matrix if row is not None]
    if not rows:
        return []
    header_index = matrix_header_index(rows)
    columns: list[tuple[int, str]] = []
    seen: set[str] = set()
    for index, cell in enumerate(rows[header_index]):
        name = cell_to_text(cell)
        if not name:
            continue
        key, suffix = name, 2
        while key in seen:
            key = f"{name} {suffix}"
            suffix += 1
        seen.add(key)
        columns.append((index, key))
    result: list[dict[str, str]] = []
    for row in rows[header_index + 1:]:
        values = [cell_to_text(cell) for cell in row]
        if all(value == "" for value in values):
            continue
        record = {name: values[index] if index < len(values) else "" for index, name in columns}
        if any(value for value in record.values()):
            result.append(record)
    return result


def records_from_matrices(matrices: list[list[list[Any]]], label: str = "archivo") -> list[dict[str, str]]:
    """Combina las hojas/tablas válidas de un libro y elige la mejor si ninguna declara encabezados."""
    best_rank = (-1, -1)
    best_records: list[dict[str, str]] = []
    combined: list[dict[str, str]] = []
    for matrix in matrices:
        if not matrix or not any(any(cell_to_text(cell) for cell in row) for row in matrix):
            continue
        header_index = matrix_header_index(matrix)
        score = header_score(matrix[header_index])
        records = matrix_to_records(matrix)
        rank = (score, len(matrix))
        if rank > best_rank:
            best_rank = rank
            best_records = records
        if score >= 2:
            combined.extend(records)
    if combined:
        return combined
    if not best_records:
        raise ValueError(f"El {label} no contiene encabezados ni filas de datos.")
    return best_records


def decode_text_bytes(file_bytes: bytes) -> str:
    encodings = ["utf-8-sig"]
    if file_bytes[:2] in {b"\xff\xfe", b"\xfe\xff"}:
        encodings.insert(0, "utf-16")
    elif b"\x00" in file_bytes:
        encodings.extend(["utf-16-le", "utf-16-be"])
    encodings.extend(["cp1252", "latin-1"])
    for encoding in encodings:
        try:
            text = file_bytes.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
        if "\x00" in text:
            continue
        visible = sum(1 for char in text if char.isprintable() or char in "\r\n\t")
        if text and visible / len(text) >= 0.85:
            return text
    raise ValueError("El archivo no parece un archivo de texto legible. Usa UTF-8, UTF-16 o Windows-1252.")


def text_delimiter(text: str) -> tuple[str, bool]:
    """Devuelve (separador, modo_espacios). No se fía de la extensión."""
    lines = [line for line in text.splitlines()[:40] if line.strip()]
    if not lines:
        return ",", False
    try:
        detected = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        if any(line.count(detected) for line in lines):
            return detected, False
    except csv.Error:
        pass
    counts = {delimiter: sum(line.count(delimiter) for line in lines) for delimiter in ",;\t|"}
    best = max(counts, key=counts.get)
    if counts[best]:
        return best, False
    whitespace_lines = sum(1 for line in lines if re.search(r"\s{2,}", line.strip()))
    if whitespace_lines >= max(2, len(lines) // 2):
        return "\x00", True
    return ",", False


def rows_from_text(file_bytes: bytes, label: str = "archivo de texto") -> list[dict[str, str]]:
    text = decode_text_bytes(file_bytes)
    delimiter, whitespace = text_delimiter(text)
    if whitespace:
        matrix = [re.split(r"\s{2,}", line.strip()) for line in text.splitlines() if line.strip()]
    else:
        matrix = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    records = matrix_to_records(matrix)
    if not records:
        raise ValueError(f"El {label} no contiene encabezados ni filas de datos.")
    return records


def rows_from_csv(file_bytes: bytes) -> list[dict[str, str]]:
    return rows_from_text(file_bytes, "archivo CSV")


def rows_from_xlsx(file_bytes: bytes) -> list[dict[str, str]]:
    import openpyxl
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("No se pudo leer el libro de Excel (.xlsx/.xlsm).") from exc
    try:
        if not workbook.worksheets:
            raise ValueError("El libro de Excel no contiene hojas.")
        matrices = [[list(row) for row in sheet.iter_rows(values_only=True)] for sheet in workbook.worksheets]
        return records_from_matrices(matrices, "libro de Excel")
    finally:
        workbook.close()


def rows_from_xls(file_bytes: bytes) -> list[dict[str, str]]:
    try:
        import xlrd
    except Exception as exc:
        raise ValueError("Para importar .xls instala xlrd (incluido en requirements.txt).") from exc
    try:
        book = xlrd.open_workbook(file_contents=file_bytes)
    except Exception as exc:
        raise ValueError("No se pudo leer el libro de Excel antiguo (.xls).") from exc
    matrices: list[list[list[Any]]] = []
    for sheet in book.sheets():
        matrix: list[list[Any]] = []
        for row_index in range(sheet.nrows):
            row: list[Any] = []
            for column_index in range(sheet.ncols):
                cell = sheet.cell(row_index, column_index)
                value: Any = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                    except Exception:
                        value = cell.value
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(cell.value)
                row.append(value)
            matrix.append(row)
        matrices.append(matrix)
    return records_from_matrices(matrices, "libro de Excel .xls")


ODS_TABLE = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
ODS_OFFICE = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
ODS_TEXT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"


def ods_cell_text(cell: ET.Element) -> str:
    paragraphs: list[str] = []
    for paragraph in cell.iter(f"{ODS_TEXT}p"):
        text = "".join(paragraph.itertext()).strip()
        if text:
            paragraphs.append(text)
    return " ".join(paragraphs)


def ods_duration(value: str) -> str:
    match = re.fullmatch(r"PT(?:(\d+)H)?(?:(\d+)M)?(?:(\d+(?:\.\d+)?)S)?", value or "")
    if not match:
        return value
    hours = int(match.group(1) or 0)
    minutes = int(match.group(2) or 0)
    seconds = float(match.group(3) or 0)
    total_minutes = hours * 60 + minutes + round(seconds / 60)
    return f"{total_minutes // 60:02d}:{total_minutes % 60:02d}"


def ods_cell_value(cell: ET.Element) -> Any:
    value_type = cell.get(f"{ODS_OFFICE}value-type")
    if value_type == "date":
        return cell.get(f"{ODS_OFFICE}date-value") or ods_cell_text(cell)
    if value_type == "time":
        return ods_duration(cell.get(f"{ODS_OFFICE}time-value") or ods_cell_text(cell))
    if value_type in {"float", "currency", "percentage"}:
        return cell.get(f"{ODS_OFFICE}value") or ods_cell_text(cell)
    if value_type == "boolean":
        return "true" if cell.get(f"{ODS_OFFICE}boolean-value") == "true" else "false"
    return ods_cell_text(cell)


def rows_from_ods(file_bytes: bytes) -> list[dict[str, str]]:
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            content = archive.read("content.xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValueError("No se pudo leer el archivo ODS (.ods).") from exc
    try:
        root = ET.fromstring(content)
    except ET.ParseError as exc:
        raise ValueError("El contenido del ODS no es XML válido.") from exc
    matrices: list[list[list[Any]]] = []
    for table in root.iter(f"{ODS_TABLE}table"):
        matrix: list[list[Any]] = []
        for row in table.iter(f"{ODS_TABLE}table-row"):
            values: list[Any] = []
            for cell in row:
                if cell.tag not in {f"{ODS_TABLE}table-cell", f"{ODS_TABLE}covered-table-cell"}:
                    continue
                repeat = int(cell.get(f"{ODS_TABLE}number-columns-repeated", "1") or 1)
                value = ods_cell_value(cell) if cell.tag == f"{ODS_TABLE}table-cell" else ""
                values.extend([value] * min(repeat, 200))
            if not any(cell_to_text(value) for value in values):
                continue
            row_repeat = int(row.get(f"{ODS_TABLE}number-rows-repeated", "1") or 1)
            for _ in range(min(row_repeat, 5000)):
                matrix.append(list(values))
            if len(matrix) > 20000:
                break
        matrices.append(matrix)
    return records_from_matrices(matrices, "archivo ODS")


def rows_from_pdf(file_bytes: bytes) -> list[dict[str, str]]:
    import fitz
    try:
        document = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:
        raise ValueError("No se pudo abrir el PDF.") from exc
    try:
        matrices: list[list[list[Any]]] = []
        text_lines: list[str] = []
        for page in document:
            for table in page.find_tables():
                matrix = table.extract()
                if matrix and any(any(cell_to_text(cell) for cell in row) for row in matrix):
                    matrices.append(matrix)
            text_lines.extend(page.get_text("text").splitlines())
        if matrices:
            try:
                return records_from_matrices(matrices, "PDF")
            except ValueError:
                pass
        fallback = "\n".join(text_lines).strip()
        if fallback:
            try:
                return rows_from_text(fallback.encode("utf-8"), "PDF")
            except ValueError:
                pass
        raise ValueError("No se detectó una tabla ni texto tabular en el PDF. Los PDF escaneados (imagen) todavía no se admiten.")
    finally:
        document.close()


def detect_import_kind(file_bytes: bytes) -> str:
    if not file_bytes or not file_bytes.strip():
        raise ValueError("El archivo está vacío.")
    if file_bytes.startswith(b"%PDF"):
        return "pdf"
    if file_bytes.startswith(b"PK"):
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
                names = set(archive.namelist())
                if "xl/workbook.xml" in names:
                    return "xlsx"
                if "content.xml" in names:
                    content = archive.read("content.xml")[:4096]
                    if b"document-content" in content or b"table:table" in content:
                        return "ods"
        except zipfile.BadZipFile as exc:
            raise ValueError("El archivo comprimido no es un Excel/ODS válido.") from exc
        raise ValueError("El archivo comprimido no parece un Excel (.xlsx/.xlsm) ni un ODS.")
    if file_bytes.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls"
    return "text"


def read_import_file(filename: str, file_bytes: bytes) -> tuple[str, list[dict[str, str]]]:
    """Detecta el formato por contenido real y devuelve filas crudas (sin escribir nada)."""
    kind = detect_import_kind(file_bytes)
    if kind == "pdf":
        return kind, rows_from_pdf(file_bytes)
    if kind == "xlsx":
        return kind, rows_from_xlsx(file_bytes)
    if kind == "xls":
        return kind, rows_from_xls(file_bytes)
    if kind == "ods":
        return kind, rows_from_ods(file_bytes)
    return "text", rows_from_text(file_bytes, "archivo de texto")


def create_app(db_path: str | Path | None = None) -> FastAPI:
    selected_path = Path(db_path or os.environ.get("SHIFTS_DB_PATH", ROOT / "data" / "quadrants.sqlite3")).resolve()
    initialize(selected_path)
    app = FastAPI(title="Turnos · Gestor local", docs_url=None, redoc_url=None)
    app.state.db_path = selected_path

    @app.middleware("http")
    async def no_store_api(request: Request, call_next):
        response = await call_next(request)
        if request.url.path.startswith("/api/"):
            response.headers["Cache-Control"] = "no-store"
        return response

    @app.get("/")
    def index() -> FileResponse:
        return FileResponse(STATIC_DIR / "index.html")

    app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

    @app.get("/api/health")
    def health() -> dict[str, str]:
        return {"status": "ok", "storage": "local"}

    @app.get("/api/employees")
    def employees(q: str = "", limit: int = 5000) -> list[dict[str, Any]]:
        limit = max(1, min(limit, 5000))
        with closing(connect(selected_path)) as conn:
            rows = conn.execute(
                "SELECT id,code,first_name,last_name,active FROM employees WHERE active=1 AND (code LIKE ? OR first_name LIKE ? OR last_name LIKE ?) ORDER BY code LIMIT ?",
                (f"%{q}%", f"%{q}%", f"%{q}%", limit),
            ).fetchall()
            return [row_dict(row) for row in rows]

    @app.post("/api/employees", status_code=201)
    def add_employee(data: EmployeeInput) -> dict[str, Any]:
        with closing(connect(selected_path)) as conn:
            existing = conn.execute("SELECT id,active FROM employees WHERE code=?", (data.code.strip(),)).fetchone()
            if existing:
                if existing["active"]:
                    raise HTTPException(409, "Ese ID de empleado ya existe.")
                conn.execute("UPDATE employees SET first_name=?,last_name=?,active=1 WHERE id=?", (data.first_name.strip(), data.last_name.strip(), existing["id"]))
                conn.commit()
                return row_dict(conn.execute("SELECT id,code,first_name,last_name,active FROM employees WHERE id=?", (existing["id"],)).fetchone())
            employee_id = str(uuid.uuid4())
            try:
                conn.execute("INSERT INTO employees(id,code,first_name,last_name,created_at) VALUES(?,?,?,?,?)", (employee_id, data.code.strip(), data.first_name.strip(), data.last_name.strip(), now_text()))
                conn.commit()
            except sqlite3.IntegrityError as exc:
                raise HTTPException(409, "Ese ID de empleado ya existe.") from exc
            return row_dict(conn.execute("SELECT id,code,first_name,last_name,active FROM employees WHERE id=?", (employee_id,)).fetchone())

    @app.patch("/api/employees/{employee_id}")
    def update_employee(employee_id: str, data: EmployeeInput) -> dict[str, Any]:
        with closing(connect(selected_path)) as conn:
            if not conn.execute("SELECT 1 FROM employees WHERE id=?", (employee_id,)).fetchone():
                raise HTTPException(404, "Empleado no encontrado.")
            conflict = conn.execute("SELECT 1 FROM employees WHERE code=? AND id<>?", (data.code.strip(), employee_id)).fetchone()
            if conflict:
                raise HTTPException(409, "Ese ID de empleado ya existe.")
            conn.execute("UPDATE employees SET code=?,first_name=?,last_name=? WHERE id=?", (data.code.strip(), data.first_name.strip(), data.last_name.strip(), employee_id))
            conn.commit()
            return row_dict(conn.execute("SELECT id,code,first_name,last_name,active FROM employees WHERE id=?", (employee_id,)).fetchone())

    @app.delete("/api/employees/{employee_id}", status_code=204)
    def delete_employee(employee_id: str) -> None:
        with closing(connect(selected_path)) as conn:
            result = conn.execute("UPDATE employees SET active=0 WHERE id=? AND active=1", (employee_id,))
            if not result.rowcount:
                raise HTTPException(404, "Empleado no encontrado.")
            conn.commit()

    @app.get("/api/posts")
    def posts(q: str = "") -> list[dict[str, Any]]:
        with closing(connect(selected_path)) as conn:
            rows = conn.execute(
                "SELECT id,code,name,municipality,province,active FROM posts WHERE active=1 AND (code LIKE ? OR name LIKE ? OR municipality LIKE ?) ORDER BY code, name",
                (f"%{q}%", f"%{q}%", f"%{q}%"),
            ).fetchall()
            return [row_dict(row) for row in rows]

    @app.post("/api/posts", status_code=201)
    def add_post(data: PostInput) -> dict[str, Any]:
        code = data.code.strip().upper()
        name = data.name.strip()
        municipality = data.municipality.strip() or DEFAULT_MUNICIPALITY
        province = data.province.strip() or DEFAULT_PROVINCE
        with closing(connect(selected_path)) as conn:
            by_code = conn.execute("SELECT id,active,code FROM posts WHERE UPPER(code)=?", (code,)).fetchone() if code else None
            by_name = conn.execute("SELECT id,active,code FROM posts WHERE name=? AND municipality=?", (name, municipality)).fetchone()
            if by_code and by_name and by_code["id"] != by_name["id"]:
                raise HTTPException(409, "El ID y el nombre pertenecen a centros distintos.")
            target = by_code or by_name
            if target:
                if target["active"]:
                    if by_code:
                        raise HTTPException(409, "Ese ID de centro ya existe.")
                    raise HTTPException(409, "Ya existe un centro con ese nombre.")
                final_code = code or target["code"] or generated_post_code(conn, name)
                conflict = conn.execute("SELECT 1 FROM posts WHERE UPPER(code)=? AND id<>?", (final_code, target["id"])).fetchone()
                if conflict:
                    raise HTTPException(409, "Ese ID de centro ya existe.")
                conn.execute("UPDATE posts SET code=?,name=?,municipality=?,province=?,active=1 WHERE id=?", (final_code, name, municipality, province, target["id"]))
                sync_official_holidays(conn)
                conn.commit()
                return row_dict(conn.execute("SELECT id,code,name,municipality,province,active FROM posts WHERE id=?", (target["id"],)).fetchone())
            post_id = str(uuid.uuid4())
            final_code = code or generated_post_code(conn, name)
            try:
                conn.execute("INSERT INTO posts(id,code,name,municipality,province,created_at) VALUES(?,?,?,?,?,?)", (post_id, final_code, name, municipality, province, now_text()))
                sync_official_holidays(conn)
                conn.commit()
            except sqlite3.IntegrityError as exc:
                raise HTTPException(409, "Ya existe un centro con ese ID o nombre.") from exc
            return row_dict(conn.execute("SELECT id,code,name,municipality,province,active FROM posts WHERE id=?", (post_id,)).fetchone())

    @app.patch("/api/posts/{post_id}")
    def update_post(post_id: str, data: PostInput) -> dict[str, Any]:
        with closing(connect(selected_path)) as conn:
            current = conn.execute("SELECT id,code,name,municipality,province FROM posts WHERE id=?", (post_id,)).fetchone()
            if not current:
                raise HTTPException(404, "Centro no encontrado.")
            code = data.code.strip().upper() or current["code"]
            name = data.name.strip()
            municipality = data.municipality.strip() or current["municipality"]
            province = data.province.strip() or current["province"]
            conflict_code = conn.execute("SELECT 1 FROM posts WHERE UPPER(code)=? AND id<>?", (code, post_id)).fetchone()
            if conflict_code:
                raise HTTPException(409, "Ese ID de centro ya existe.")
            conflict = conn.execute("SELECT 1 FROM posts WHERE name=? AND municipality=? AND id<>?", (name, municipality, post_id)).fetchone()
            if conflict:
                raise HTTPException(409, "Ya existe otro centro con ese nombre.")
            conn.execute("UPDATE posts SET code=?,name=?,municipality=?,province=? WHERE id=?", (code, name, municipality, province, post_id))
            sync_official_holidays(conn)
            conn.commit()
            return row_dict(conn.execute("SELECT id,code,name,municipality,province,active FROM posts WHERE id=?", (post_id,)).fetchone())

    @app.delete("/api/posts/{post_id}", status_code=204)
    def delete_post(post_id: str) -> None:
        with closing(connect(selected_path)) as conn:
            result = conn.execute("UPDATE posts SET active=0 WHERE id=? AND active=1", (post_id,))
            if not result.rowcount:
                raise HTTPException(404, "Centro no encontrado.")
            conn.commit()

    @app.get("/api/shifts")
    def shifts(from_date: str | None = None, to_date: str | None = None, status: str | None = None) -> list[dict[str, Any]]:
        clauses: list[str] = []
        params: list[Any] = []
        if from_date:
            clauses.append("s.end_at >= ?")
            params.append(from_date)
        if to_date:
            clauses.append("s.start_at < ?")
            params.append(to_date)
        if status in {"pending", "confirmed"}:
            clauses.append("s.status = ?")
            params.append(status)
        where = " WHERE " + " AND ".join(clauses) if clauses else ""
        with closing(connect(selected_path)) as conn:
            rows = conn.execute(
                "SELECT s.*, e.code employee_code, e.first_name, e.last_name, p.code post_code, p.name post_name, p.municipality, p.province FROM shifts s JOIN employees e ON e.id=s.employee_id JOIN posts p ON p.id=s.post_id" + where + " ORDER BY s.start_at LIMIT 20000",
                params,
            ).fetchall()
            return [row_dict(row) for row in rows]

    @app.post("/api/shifts", status_code=201)
    def add_shift(data: ShiftInput) -> dict[str, Any]:
        mode = (data.repeat_mode or "none").strip().lower()
        if mode not in {"none", "daily", "weekdays", "weekly", "custom"}:
            raise HTTPException(422, "Opción de repetición no válida.")
        with closing(connect(selected_path)) as conn:
            try:
                start, end, break_start, break_end, second_start, second_end = validate_shift(data, conn)
                if mode == "none":
                    days = [start.date()]
                elif mode == "custom":
                    raw_dates = data.repeat_dates or []
                    if not raw_dates:
                        raise ValueError("Selecciona al menos un día concreto para la repetición esporádica.")
                    days = sorted({coerce_date(value) for value in raw_dates})
                    if len(days) > 366:
                        raise ValueError("La repetición admite como máximo 366 días.")
                    earlier = [day for day in days if day < start.date()]
                    if earlier:
                        raise ValueError("Los días concretos deben ser iguales o posteriores a la entrada.")
                else:
                    if not data.repeat_until:
                        raise ValueError("Indica hasta qué fecha repetir el turno.")
                    until = coerce_date(data.repeat_until)
                    if until < start.date():
                        raise ValueError("La fecha final de repetición debe ser igual o posterior a la entrada.")
                    if (until - start.date()).days > 366:
                        raise ValueError("La repetición admite como máximo 366 días.")
                    days = []
                    current = start.date()
                    while current <= until:
                        if mode == "daily" or (mode == "weekdays" and current.weekday() < 5) or (mode == "weekly" and current.weekday() == start.date().weekday()):
                            days.append(current)
                        current += timedelta(days=1)
                    if not days:
                        raise ValueError("La repetición seleccionada no genera ningún día en ese rango.")
            except ValueError as exc:
                raise HTTPException(422, str(exc)) from exc
            created: list[str] = []
            skipped = 0
            for day in days:
                offset = day - start.date()
                occurrence = ShiftInput(
                    employee_id=data.employee_id,
                    post_id=data.post_id,
                    start_at=(start + offset).isoformat(timespec="minutes"),
                    end_at=(end + offset).isoformat(timespec="minutes"),
                    second_start_at=(second_start + offset).isoformat(timespec="minutes") if second_start else None,
                    second_end_at=(second_end + offset).isoformat(timespec="minutes") if second_end else None,
                    break_start_at=(break_start + offset).isoformat(timespec="minutes") if break_start else None,
                    break_end_at=(break_end + offset).isoformat(timespec="minutes") if break_end else None,
                    break_paid=data.break_paid,
                )
                try:
                    occ_start, occ_end, occ_break_start, occ_break_end, occ_second_start, occ_second_end = validate_shift(occurrence, conn)
                except ValueError as exc:
                    raise HTTPException(422, f"{day.isoformat()}: {exc}") from exc
                start_text = occ_start.isoformat(timespec="minutes")
                end_text = occ_end.isoformat(timespec="minutes")
                break_start_text = occ_break_start.isoformat(timespec="minutes") if occ_break_start else None
                break_end_text = occ_break_end.isoformat(timespec="minutes") if occ_break_end else None
                second_start_text = occ_second_start.isoformat(timespec="minutes") if occ_second_start else None
                second_end_text = occ_second_end.isoformat(timespec="minutes") if occ_second_end else None
                duplicate = conn.execute(
                    "SELECT 1 FROM shifts WHERE employee_id=? AND post_id=? AND start_at=? AND end_at=? AND COALESCE(second_start_at,'')=? AND COALESCE(second_end_at,'')=? AND COALESCE(break_start_at,'')=? AND COALESCE(break_end_at,'')=? AND break_paid=? AND status IN ('confirmed','pending')",
                    (data.employee_id, data.post_id, start_text, end_text, second_start_text or "", second_end_text or "", break_start_text or "", break_end_text or "", int(bool(data.break_paid))),
                ).fetchone()
                if duplicate:
                    skipped += 1
                    continue
                shift_id = str(uuid.uuid4())
                conn.execute(
                    "INSERT INTO shifts(id,employee_id,post_id,start_at,end_at,second_start_at,second_end_at,break_start_at,break_end_at,break_paid,status,source,created_at) VALUES(?,?,?,?,?,?,?,?,?,?, 'confirmed','manual',?)",
                    (shift_id, data.employee_id, data.post_id, start_text, end_text, second_start_text, second_end_text, break_start_text, break_end_text, int(bool(data.break_paid)), now_text()),
                )
                created.append(shift_id)
            conn.commit()
            return {"id": created[0] if created else "", "status": "confirmed", "created": len(created), "skipped": skipped}

    @app.patch("/api/shifts/{shift_id}")
    def update_shift(shift_id: str, data: ShiftInput) -> dict[str, str]:
        with closing(connect(selected_path)) as conn:
            current = conn.execute("SELECT status FROM shifts WHERE id=?", (shift_id,)).fetchone()
            if not current:
                raise HTTPException(404, "No se encontró el turno.")
            try:
                start, end, break_start, break_end, second_start, second_end = validate_shift(data, conn)
                conn.execute(
                    "UPDATE shifts SET employee_id=?,post_id=?,start_at=?,end_at=?,second_start_at=?,second_end_at=?,break_start_at=?,break_end_at=?,break_paid=? WHERE id=?",
                    (data.employee_id, data.post_id, start.isoformat(timespec="minutes"), end.isoformat(timespec="minutes"), second_start.isoformat(timespec="minutes") if second_start else None, second_end.isoformat(timespec="minutes") if second_end else None, break_start.isoformat(timespec="minutes") if break_start else None, break_end.isoformat(timespec="minutes") if break_end else None, int(bool(data.break_paid)), shift_id),
                )
                conn.commit()
            except ValueError as exc:
                raise HTTPException(422, str(exc)) from exc
            return {"id": shift_id, "status": current["status"]}

    @app.post("/api/shifts/{shift_id}/confirm")
    def confirm_shift(shift_id: str) -> dict[str, str]:
        with closing(connect(selected_path)) as conn:
            result = conn.execute("UPDATE shifts SET status='confirmed' WHERE id=? AND status='pending'", (shift_id,))
            if not result.rowcount:
                raise HTTPException(404, "No se encontró un turno pendiente.")
            conn.commit()
            return {"status": "confirmed"}

    @app.post("/api/shifts/review")
    def send_shifts_to_review(data: ShiftIds) -> dict[str, int]:
        ids = list(dict.fromkeys(data.ids))
        with closing(connect(selected_path)) as conn:
            placeholders = ",".join("?" for _ in ids)
            result = conn.execute(f"UPDATE shifts SET status='pending' WHERE status='confirmed' AND id IN ({placeholders})", ids)
            conn.commit()
            return {"updated": result.rowcount, "requested": len(ids)}

    @app.post("/api/review/{review_id}/resolve", status_code=201)
    def resolve_review(review_id: str, data: ShiftInput) -> dict[str, str]:
        with closing(connect(selected_path)) as conn:
            if not conn.execute("SELECT 1 FROM review_items WHERE id=?", (review_id,)).fetchone():
                raise HTTPException(404, "El elemento de revisión ya no existe.")
            try:
                start, end, break_start, break_end, second_start, second_end = validate_shift(data, conn)
                shift_id = str(uuid.uuid4())
                conn.execute(
                    "INSERT INTO shifts(id,employee_id,post_id,start_at,end_at,second_start_at,second_end_at,break_start_at,break_end_at,break_paid,status,source,source_label,created_at) VALUES(?,?,?,?,?,?,?,?,?,?, 'confirmed','csv','Revisado manualmente',?)",
                    (shift_id, data.employee_id, data.post_id, start.isoformat(timespec="minutes"), end.isoformat(timespec="minutes"), second_start.isoformat(timespec="minutes") if second_start else None, second_end.isoformat(timespec="minutes") if second_end else None, break_start.isoformat(timespec="minutes") if break_start else None, break_end.isoformat(timespec="minutes") if break_end else None, int(bool(data.break_paid)), now_text()),
                )
                conn.execute("DELETE FROM review_items WHERE id=?", (review_id,))
                conn.commit()
                return {"id": shift_id, "status": "confirmed"}
            except ValueError as exc:
                raise HTTPException(422, str(exc)) from exc

    @app.post("/api/review/{review_id}/resolve-summary", status_code=201)
    def resolve_summary_review(review_id: str, data: SummaryReviewInput) -> dict[str, Any]:
        with closing(connect(selected_path)) as conn:
            item = conn.execute("SELECT * FROM review_items WHERE id=?", (review_id,)).fetchone()
            if not item:
                raise HTTPException(404, "El elemento de revisión ya no existe.")
            if not conn.execute("SELECT 1 FROM employees WHERE id=? AND active=1", (data.employee_id,)).fetchone():
                raise HTTPException(422, "El empleado seleccionado no existe.")
            try:
                month = parse_month(data.month)
                minutes = parse_hours(data.hours)
            except ValueError as exc:
                raise HTTPException(422, str(exc)) from exc
            raw = json.loads(item["raw_json"])
            post_id = raw.get("post_id") or None
            if not post_id:
                post_id = resolve_or_create_post(conn, raw) or None
            now = now_text()
            existing = conn.execute("SELECT id FROM monthly_summaries WHERE employee_id=? AND month=?", (data.employee_id, month)).fetchone()
            if existing:
                conn.execute("UPDATE monthly_summaries SET post_id=?, minutes=?, source=?, source_label=?, created_at=? WHERE id=?", (post_id, minutes, "manual", "Revisado manualmente", now, existing["id"]))
                summary_id = existing["id"]
            else:
                summary_id = str(uuid.uuid4())
                conn.execute("INSERT INTO monthly_summaries(id,employee_id,post_id,month,minutes,source,source_label,created_at) VALUES(?,?,?,?,?,?,?,?)", (summary_id, data.employee_id, post_id, month, minutes, "manual", "Revisado manualmente", now))
            conn.execute("DELETE FROM review_items WHERE id=?", (review_id,))
            conn.commit()
            return {"id": summary_id, "month": month, "minutes": minutes, "status": "saved"}

    @app.delete("/api/review/{review_id}", status_code=204)
    def discard_review(review_id: str) -> None:
        with closing(connect(selected_path)) as conn:
            result = conn.execute("DELETE FROM review_items WHERE id=?", (review_id,))
            if not result.rowcount:
                raise HTTPException(404, "Elemento de revisión no encontrado.")
            conn.commit()

    @app.delete("/api/shifts/{shift_id}", status_code=204)
    def delete_shift(shift_id: str) -> None:
        with closing(connect(selected_path)) as conn:
            result = conn.execute("DELETE FROM shifts WHERE id=? AND status='pending'", (shift_id,))
            if not result.rowcount:
                raise HTTPException(404, "Solo se pueden descartar turnos pendientes.")
            conn.commit()

    @app.get("/api/holidays")
    def holidays(year: int | None = None, municipality: str | None = None) -> list[dict[str, Any]]:
        clauses: list[str] = []
        params: list[Any] = []
        if year:
            clauses.append("date >= ? AND date < ?")
            params.extend([f"{year:04d}-01-01", f"{year + 1:04d}-01-01"])
        if municipality:
            clauses.append("municipality = ?")
            params.append(municipality)
        where = " WHERE " + " AND ".join(clauses) if clauses else ""
        with closing(connect(selected_path)) as conn:
            return [row_dict(row) for row in conn.execute("SELECT * FROM holidays" + where + " ORDER BY date", params).fetchall()]

    @app.post("/api/holidays", status_code=201)
    def add_holiday(data: HolidayInput) -> dict[str, Any]:
        try:
            parsed_date = date.fromisoformat(data.date)
        except ValueError as exc:
            raise HTTPException(422, "Fecha festiva no válida.") from exc
        holiday_id = str(uuid.uuid4())
        with closing(connect(selected_path)) as conn:
            try:
                conn.execute("INSERT INTO holidays(id,date,municipality,name,created_at) VALUES(?,?,?,?,?)", (holiday_id, parsed_date.isoformat(), data.municipality.strip(), data.name.strip(), now_text()))
                conn.execute("DELETE FROM dismissed_holidays WHERE date=? AND municipality=? AND name=?", (parsed_date.isoformat(), data.municipality.strip(), data.name.strip()))
                conn.commit()
            except sqlite3.IntegrityError as exc:
                raise HTTPException(409, "Ese festivo ya está registrado.") from exc
            return row_dict(conn.execute("SELECT * FROM holidays WHERE id=?", (holiday_id,)).fetchone())

    @app.delete("/api/holidays/{holiday_id}", status_code=204)
    def delete_holiday(holiday_id: str) -> None:
        with closing(connect(selected_path)) as conn:
            row = conn.execute("SELECT date,municipality,name FROM holidays WHERE id=?", (holiday_id,)).fetchone()
            if not row:
                raise HTTPException(404, "Festivo no encontrado.")
            conn.execute("INSERT OR IGNORE INTO dismissed_holidays(date,municipality,name) VALUES(?,?,?)", (row["date"], row["municipality"], row["name"]))
            conn.execute("DELETE FROM holidays WHERE id=?", (holiday_id,))
            conn.commit()

    @app.get("/api/monthly-summaries")
    def monthly_summaries(year: int | None = None, month: int | None = None) -> list[dict[str, Any]]:
        clauses: list[str] = []
        params: list[Any] = []
        if year:
            if month:
                clauses.append("ms.month=?")
                params.append(f"{year:04d}-{month:02d}")
            else:
                clauses.append("ms.month LIKE ?")
                params.append(f"{year:04d}-%")
        where = " WHERE " + " AND ".join(clauses) if clauses else ""
        with closing(connect(selected_path)) as conn:
            rows = conn.execute("SELECT ms.*, e.code employee_code, e.first_name, e.last_name, p.code post_code, p.name post_name FROM monthly_summaries ms JOIN employees e ON e.id=ms.employee_id LEFT JOIN posts p ON p.id=ms.post_id" + where + " ORDER BY ms.month, e.code", params).fetchall()
            return [{"id": r["id"], "employee_code": r["employee_code"], "name": f"{r['first_name']} {r['last_name']}".strip(), "post_code": r["post_code"], "post_name": r["post_name"], "month": r["month"], "minutes": r["minutes"], "hours": hhmm(r["minutes"]), "source_label": r["source_label"]} for r in rows]

    @app.delete("/api/monthly-summaries/{summary_id}", status_code=204)
    def delete_summary(summary_id: str) -> None:
        with closing(connect(selected_path)) as conn:
            result = conn.execute("DELETE FROM monthly_summaries WHERE id=?", (summary_id,))
            if not result.rowcount:
                raise HTTPException(404, "Resumen mensual no encontrado.")
            conn.commit()

    @app.get("/api/duplicates")
    def duplicates() -> list[dict[str, Any]]:
        """Empleados activos con el mismo nombre y distinto ID (para revisar/unificar)."""
        with closing(connect(selected_path)) as conn:
            groups: dict[str, list[dict[str, Any]]] = {}
            for row in conn.execute("SELECT id,code,first_name,last_name FROM employees WHERE active=1 ORDER BY code").fetchall():
                key = normalized_name(row["first_name"], row["last_name"])
                if not key:
                    continue
                groups.setdefault(key, []).append({"id": row["id"], "code": row["code"], "name": f"{row['first_name']} {row['last_name']}".strip()})
            result: list[dict[str, Any]] = []
            for people in groups.values():
                if len(people) < 2:
                    continue
                for person in people:
                    person["shift_count"] = conn.execute("SELECT COUNT(*) FROM shifts WHERE employee_id=?", (person["id"],)).fetchone()[0]
                    person["summary_count"] = conn.execute("SELECT COUNT(*) FROM monthly_summaries WHERE employee_id=?", (person["id"],)).fetchone()[0]
                result.append({"name": people[0]["name"], "employees": people})
            return result

    @app.post("/api/employees/{target_id}/merge")
    def merge_employee(target_id: str, data: MergeInput) -> dict[str, Any]:
        """Mueve turnos y resúmenes del empleado origen al destino y elimina el origen."""
        with closing(connect(selected_path)) as conn:
            if target_id == data.source_id:
                raise HTTPException(422, "Origen y destino son el mismo empleado.")
            if not conn.execute("SELECT 1 FROM employees WHERE id=?", (target_id,)).fetchone():
                raise HTTPException(404, "Empleado destino no encontrado.")
            if not conn.execute("SELECT 1 FROM employees WHERE id=?", (data.source_id,)).fetchone():
                raise HTTPException(404, "Empleado origen no encontrado.")
            moved_shifts = conn.execute("UPDATE shifts SET employee_id=? WHERE employee_id=?", (target_id, data.source_id)).rowcount
            moved_summaries = 0
            for row in conn.execute("SELECT id,month,minutes FROM monthly_summaries WHERE employee_id=?", (data.source_id,)).fetchall():
                existing = conn.execute("SELECT id FROM monthly_summaries WHERE employee_id=? AND month=?", (target_id, row["month"])).fetchone()
                if existing:
                    conn.execute("UPDATE monthly_summaries SET minutes=minutes+? WHERE id=?", (row["minutes"], existing["id"]))
                    conn.execute("DELETE FROM monthly_summaries WHERE id=?", (row["id"],))
                else:
                    conn.execute("UPDATE monthly_summaries SET employee_id=? WHERE id=?", (target_id, row["id"]))
                moved_summaries += 1
            conn.execute("DELETE FROM employees WHERE id=?", (data.source_id,))
            conn.commit()
            return {"merged_shifts": moved_shifts, "merged_summaries": moved_summaries, "target_id": target_id}

    @app.get("/api/reports")
    def reports(year: int, month: int | None = None, q: str = "", post_id: str = "", group_by_employee: bool = True) -> dict[str, Any]:
        if year < 2000 or year > 2200 or (month is not None and not 1 <= month <= 12):
            raise HTTPException(422, "Selecciona un año y mes válidos.")
        first_day = date(year, month or 1, 1)
        if month:
            last_day = date(year + (month == 12), 1 if month == 12 else month + 1, 1)
        else:
            last_day = date(year + 1, 1, 1)
        start_bound = datetime.combine(first_day, time.min, LOCAL_ZONE)
        end_bound = datetime.combine(last_day, time.min, LOCAL_ZONE)
        with closing(connect(selected_path)) as conn:
            clauses = ["s.status='confirmed'", "s.start_at < ?", "s.end_at > ?", "(e.code LIKE ? OR e.first_name LIKE ? OR e.last_name LIKE ?)"]
            params: list[Any] = [end_bound.replace(tzinfo=None).isoformat(timespec="minutes"), start_bound.replace(tzinfo=None).isoformat(timespec="minutes"), f"%{q}%", f"%{q}%", f"%{q}%"]
            if post_id:
                clauses.append("s.post_id=?")
                params.append(post_id)
            selected = conn.execute(
                "SELECT s.*, e.code employee_code,e.first_name,e.last_name,p.code post_code,p.name post_name,p.municipality,p.province FROM shifts s JOIN employees e ON e.id=s.employee_id JOIN posts p ON p.id=s.post_id WHERE " + " AND ".join(clauses) + " ORDER BY s.start_at",
                params,
            ).fetchall()
            municipalities = {row["municipality"] for row in selected}
            holiday_dates: dict[str, set[str]] = {municipality: set() for municipality in municipalities}
            for h in conn.execute("SELECT date,municipality FROM holidays WHERE date>=? AND date<?", (first_day.isoformat(), last_day.isoformat())).fetchall():
                holiday_dates.setdefault(h["municipality"], set()).add(h["date"])
            summary_clauses = ["ms.month=?" if month else "ms.month LIKE ?", "(e.code LIKE ? OR e.first_name LIKE ? OR e.last_name LIKE ?)"]
            summary_params: list[Any] = [f"{year:04d}-{month:02d}" if month else f"{year:04d}-%", f"%{q}%", f"%{q}%", f"%{q}%"]
            if post_id:
                summary_clauses.append("ms.post_id=?")
                summary_params.append(post_id)
            summary_rows = conn.execute(
                "SELECT ms.*, e.code employee_code, e.first_name, e.last_name, p.code post_code, p.name post_name FROM monthly_summaries ms JOIN employees e ON e.id=ms.employee_id LEFT JOIN posts p ON p.id=ms.post_id WHERE " + " AND ".join(summary_clauses) + " ORDER BY ms.month, e.code",
                summary_params,
            ).fetchall()
            summaries = [
                {"id": row["id"], "employee_code": row["employee_code"], "name": f"{row['first_name']} {row['last_name']}".strip(), "post_code": row["post_code"], "post_name": row["post_name"], "month": row["month"], "minutes": row["minutes"], "hours": hhmm(row["minutes"]), "source_label": row["source_label"]}
                for row in summary_rows
            ]
            by_day: dict[str, dict[str, int]] = {}
            totals = {"minutes": 0, "night_minutes": 0, "holiday_minutes": 0, "weekend_minutes": 0}
            by_employee: dict[str, dict[str, Any]] = {}
            by_post: dict[str, dict[str, Any]] = {}
            details: list[dict[str, Any]] = []
            for shift in selected:
                metrics = shift_metrics(shift, holiday_dates.get(shift["municipality"], set()), start_bound, end_bound)
                for key in totals:
                    totals[key] += metrics[key]
                if not group_by_employee:
                    for day, values in metrics["days"].items():
                        bucket = by_day.setdefault(day, {"minutes": 0, "night_minutes": 0, "holiday_minutes": 0, "shift_count": 0})
                        for key in ("minutes", "night_minutes", "holiday_minutes"):
                            bucket[key] += values[key]
                        if day == shift["start_at"][:10]:
                            bucket["shift_count"] += 1
                employee = by_employee.setdefault(shift["employee_id"], {"employee_id": shift["employee_id"], "code": shift["employee_code"], "name": f"{shift['first_name']} {shift['last_name']}".strip(), "minutes": 0, "night_minutes": 0, "holiday_minutes": 0, "shift_count": 0})
                post = by_post.setdefault(shift["post_id"], {"post_id": shift["post_id"], "code": shift["post_code"], "name": shift["post_name"], "municipality": shift["municipality"], "minutes": 0, "night_minutes": 0, "holiday_minutes": 0, "shift_count": 0})
                for key in ("minutes", "night_minutes", "holiday_minutes"):
                    employee[key] += metrics[key]
                    post[key] += metrics[key]
                employee["shift_count"] += 1
                post["shift_count"] += 1
                if group_by_employee:
                    details.append({**row_dict(shift), **metrics})
            for group in (by_employee, by_post):
                for item in group.values():
                    item["hours"] = hhmm(item["minutes"])
                    item["night_hours"] = hhmm(item["night_minutes"])
                    item["holiday_hours"] = hhmm(item["holiday_minutes"])
            return {
                "year": year, "month": month,
                "period": {"from": first_day.isoformat(), "to_exclusive": last_day.isoformat()},
                "totals": {**totals, "hours": hhmm(totals["minutes"]), "night_hours": hhmm(totals["night_minutes"]), "holiday_hours": hhmm(totals["holiday_minutes"]), "weekend_hours": hhmm(totals["weekend_minutes"])},
                "employees": sorted(by_employee.values(), key=lambda item: item["code"]),
                "posts": sorted(by_post.values(), key=lambda item: item["name"]),
                "shifts": details,
                "summaries": summaries,
                "days": [{"date": day, **values, "hours": hhmm(values["minutes"]), "night_hours": hhmm(values["night_minutes"]), "holiday_hours": hhmm(values["holiday_minutes"])} for day, values in sorted(by_day.items())] if not group_by_employee else [],
            }

    @app.get("/api/review")
    def review_queue() -> list[dict[str, Any]]:
        with closing(connect(selected_path)) as conn:
            rows = conn.execute("SELECT * FROM review_items ORDER BY created_at DESC").fetchall()
            result = [{"kind": "raw", "id": row["id"], "source_label": row["source_label"], "created_at": row["created_at"], "raw": json.loads(row["raw_json"]), "errors": json.loads(row["errors_json"])} for row in rows]
            pending = conn.execute("SELECT s.*,e.code employee_code,e.first_name,e.last_name,p.code post_code,p.name post_name,p.municipality,p.province FROM shifts s JOIN employees e ON e.id=s.employee_id JOIN posts p ON p.id=s.post_id WHERE s.status='pending' ORDER BY s.created_at DESC").fetchall()
            result.extend({"kind": "shift", "id": row["id"], "created_at": row["created_at"], "shift": row_dict(row)} for row in pending)
            return sorted(result, key=lambda item: item["created_at"], reverse=True)

    async def preview_upload(file: UploadFile) -> dict[str, Any]:
        filename = Path(file.filename or "").name
        if not filename:
            raise HTTPException(415, "Selecciona un archivo para importar.")
        content = await file.read(MAX_IMPORT_BYTES + 1)
        if len(content) > MAX_IMPORT_BYTES:
            raise HTTPException(413, "El archivo supera el límite local de 10 MB.")
        with closing(connect(selected_path)) as conn:
            try:
                kind, raw_rows = read_import_file(filename, content)
                rows = normalize_rows(raw_rows, conn)
            except ValueError as exc:
                raise HTTPException(422, str(exc)) from exc
        if not rows:
            raise HTTPException(422, "El archivo no contiene filas de turnos.")
        if len(rows) > 5000:
            raise HTTPException(422, "El archivo supera el máximo de 5000 filas. Divídelo en varios archivos.")
        return {"filename": filename, "kind": kind, "count": len(rows), "ready": sum(row["status"] == "ready" for row in rows), "pending": sum(row["status"] == "pending" for row in rows), "rows": rows}

    @app.post("/api/imports/preview")
    async def preview_import(file: UploadFile = File(...)) -> dict[str, Any]:
        return await preview_upload(file)

    @app.post("/api/imports/csv/preview")
    async def preview_csv(file: UploadFile = File(...)) -> dict[str, Any]:
        return await preview_upload(file)

    @app.post("/api/imports/confirm", status_code=201)
    def commit_import(data: ConfirmImport) -> dict[str, Any]:
        source = data.source if data.source in {"csv", "xlsx", "xls", "ods", "pdf", "text"} else "csv"
        with closing(connect(selected_path)) as conn:
            inserted = 0
            summaries = 0
            queued = 0
            now = now_text()
            try:
                for index, row in enumerate(data.rows, start=1):
                    errors = list(row.get("errors") or [])
                    label = f"{Path(data.filename).name} · fila {row.get('line', index)}"
                    if row.get("kind") == "summary":
                        summaries += store_summary_row(conn, row, errors, label, source, data.filename, now)
                        queued += 1 if errors else 0
                        continue
                    if errors:
                        queue_review(conn, label, row, errors, now)
                        queued += 1
                        continue
                    employee_id = resolve_or_create_employee(conn, row)
                    post_id = resolve_or_create_post(conn, row)
                    if not employee_id:
                        errors.append("No se pudo identificar o crear el empleado (falta el ID).")
                    if not post_id:
                        errors.append("No se pudo identificar o crear el centro (falta el nombre).")
                    shift_data = ShiftInput(
                        employee_id=employee_id, post_id=post_id,
                        start_at=str(row.get("start_at", "")), end_at=str(row.get("end_at", "")),
                        second_start_at=row.get("second_start_at"), second_end_at=row.get("second_end_at"),
                        break_start_at=row.get("break_start_at"), break_end_at=row.get("break_end_at"), break_paid=row.get("break_paid"),
                    )
                    try:
                        start, end, break_start, break_end, second_start, second_end = validate_shift(shift_data, conn) if employee_id and post_id else (None, None, None, None, None, None)
                    except ValueError as exc:
                        start = end = break_start = break_end = second_start = second_end = None
                        errors.append(str(exc))
                    if errors or not start or not end:
                        queue_review(conn, label, row, errors, now)
                        queued += 1
                        continue
                    conn.execute(
                        "INSERT INTO shifts(id,employee_id,post_id,start_at,end_at,second_start_at,second_end_at,break_start_at,break_end_at,break_paid,status,source,source_label,created_at) VALUES(?,?,?,?,?,?,?,?,?,?, 'pending',?,?,?)",
                        (str(uuid.uuid4()), shift_data.employee_id, shift_data.post_id, start.isoformat(timespec="minutes"), end.isoformat(timespec="minutes"), second_start.isoformat(timespec="minutes") if second_start else None, second_end.isoformat(timespec="minutes") if second_end else None, break_start.isoformat(timespec="minutes") if break_start else None, break_end.isoformat(timespec="minutes") if break_end else None, int(bool(shift_data.break_paid)), source, Path(data.filename).name, now),
                    )
                    inserted += 1
                conn.commit()
            except Exception:
                conn.rollback()
                raise
        return {"pending": inserted, "summaries": summaries, "needs_review": queued, "message": "Los turnos válidos se guardaron en revisión; los resúmenes mensuales se guardaron aparte; las filas ambiguas quedaron para revisar. Nada entra en los totales hasta confirmar."}

    @app.post("/api/imports/csv/confirm", status_code=201)
    def commit_csv(data: ConfirmImport) -> dict[str, Any]:
        return commit_import(data)

    @app.get("/api/exports/shifts.csv")
    def export_shifts(year: int, month: int | None = None, post_id: str = "") -> FileResponse:
        if year < 2000 or year > 2200 or (month is not None and not 1 <= month <= 12):
            raise HTTPException(422, "Periodo no válido.")
        report = reports(year=year, month=month, q="", post_id=post_id)
        output = io.StringIO(newline="")
        writer = csv.writer(output, delimiter=";")
        writer.writerow(["ID empleado","Empleado","ID centro","Centro","Entrada","Salida","Espacio libre desde","Espacio libre hasta","Horas trabajadas HH:MM","Horas nocturnas HH:MM","Horas festivas HH:MM"])
        for shift in report["shifts"]:
            free_start = shift["break_start_at"] if shift["break_start_at"] and not shift["break_paid"] else ""
            free_end = shift["break_end_at"] if shift["break_end_at"] and not shift["break_paid"] else ""
            writer.writerow([safe_export_text(shift["employee_code"]), safe_export_text(f"{shift['first_name']} {shift['last_name']}".strip()), safe_export_text(shift.get("post_code", "")), safe_export_text(shift["post_name"]), shift["start_at"], shift["end_at"], free_start, free_end, hhmm(shift["minutes"]), hhmm(shift["night_minutes"]), hhmm(shift["holiday_minutes"])])
        output.seek(0)
        from fastapi.responses import StreamingResponse
        filename = f"cuadrante-{year}-{month:02d}.csv" if month else f"cuadrante-{year}.csv"
        return StreamingResponse(iter(["\ufeff" + output.getvalue()]), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": f"attachment; filename={filename}"})

    @app.get("/api/exports/shifts.xlsx")
    def export_shifts_xlsx(year: int, month: int | None = None, post_id: str = "") -> Response:
        if year < 2000 or year > 2200 or (month is not None and not 1 <= month <= 12):
            raise HTTPException(422, "Periodo no válido.")
        report = reports(year=year, month=month, q="", post_id=post_id)
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except Exception as exc:
            raise HTTPException(500, "Falta openpyxl para exportar Excel. Instala las dependencias del proyecto.") from exc
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Cuadrante"
        headers = [
            "ID empleado", "Empleado", "ID centro", "Centro", "Fecha",
            "Entrada", "Salida", "Espacio libre desde", "Espacio libre hasta",
            "Horas trabajadas HH:MM", "Horas nocturnas HH:MM", "Horas festivas HH:MM",
        ]
        sheet.append(headers)
        for shift in report["shifts"]:
            free_start = shift["break_start_at"] if shift["break_start_at"] and not shift["break_paid"] else ""
            free_end = shift["break_end_at"] if shift["break_end_at"] and not shift["break_paid"] else ""
            sheet.append([
                safe_export_text(shift["employee_code"]),
                safe_export_text(f"{shift['first_name']} {shift['last_name']}".strip()),
                safe_export_text(shift.get("post_code", "")),
                safe_export_text(shift["post_name"]),
                shift["start_at"][:10],
                shift["start_at"][11:16],
                shift["end_at"][11:16],
                free_start[11:16] if free_start else "",
                free_end[11:16] if free_end else "",
                hhmm(shift["minutes"]),
                hhmm(shift["night_minutes"]),
                hhmm(shift["holiday_minutes"]),
            ])
        header_fill = PatternFill("solid", fgColor="19344B")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, width in enumerate([14, 28, 16, 28, 12, 10, 10, 16, 16, 20, 20, 18], start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        buffer = io.BytesIO()
        workbook.save(buffer)
        filename = f"cuadrante-{year}-{month:02d}.xlsx" if month else f"cuadrante-{year}.xlsx"
        return Response(
            content=buffer.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    return app


app = create_app()
