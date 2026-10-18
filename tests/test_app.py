from __future__ import annotations

import csv
import io
from datetime import date, datetime, time, timedelta

import pytest
from fastapi.testclient import TestClient

from main import create_app, connect, shift_metrics


@pytest.fixture()
def client(tmp_path):
    app = create_app(tmp_path / "test.sqlite3")
    with TestClient(app) as test_client:
        yield test_client, app.state.db_path


def setup_people(client):
    test_client, _ = client
    employee = test_client.post("/api/employees", json={"code": "VS-001", "first_name": "Ana", "last_name": "Vigilante"})
    post = test_client.post("/api/posts", json={"name": "Puesto Central", "municipality": "València", "province": "Valencia"})
    assert employee.status_code == 201
    assert post.status_code == 201
    return employee.json(), post.json()


def add_shift(test_client, employee, post, **overrides):
    data = {
        "employee_id": employee["id"],
        "post_id": post["id"],
        "start_at": "2026-04-10T21:30",
        "end_at": "2026-04-11T06:30",
        "break_start_at": "2026-04-11T01:00",
        "break_end_at": "2026-04-11T01:30",
        "break_paid": False,
    }
    data.update(overrides)
    return test_client.post("/api/shifts", json=data)


def test_manual_shift_night_break_and_month_aggregation(client):
    test_client, _ = client
    employee, post = setup_people(client)
    shift = add_shift(test_client, employee, post)
    assert shift.status_code == 201, shift.text
    report = test_client.get("/api/reports", params={"year": 2026, "month": 4}).json()
    assert report["totals"]["hours"] == "08:30"
    assert report["totals"]["night_hours"] == "07:30"
    assert report["employees"][0]["code"] == "VS-001"
    assert report["shifts"][0]["minutes"] == 510


def test_holiday_locality_and_overnight_segments(client):
    test_client, _ = client
    employee, post = setup_people(client)
    holiday = test_client.post("/api/holidays", json={"date": "2026-04-11", "municipality": "València", "name": "Festivo de prueba"})
    assert holiday.status_code == 201
    shift = add_shift(test_client, employee, post, start_at="2026-04-10T23:00", end_at="2026-04-11T02:00", break_start_at=None, break_end_at=None)
    assert shift.status_code == 201
    report = test_client.get("/api/reports", params={"year": 2026, "month": 4}).json()
    assert report["totals"]["hours"] == "03:00"
    assert report["totals"]["night_hours"] == "03:00"
    assert report["totals"]["holiday_hours"] == "02:00"
    assert report["totals"]["weekend_hours"] == "02:00"


def test_paid_pause_is_not_subtracted_and_equal_times_rejected(client):
    test_client, _ = client
    employee, post = setup_people(client)
    result = add_shift(test_client, employee, post, start_at="2026-04-10T21:30", end_at="2026-04-11T06:30", break_paid=True)
    assert result.status_code == 201, result.text
    report = test_client.get("/api/reports", params={"year": 2026, "month": 4}).json()
    assert report["totals"]["hours"] == "09:00"
    invalid = add_shift(test_client, employee, post, start_at="2026-04-10T08:00", end_at="2026-04-10T08:00", break_start_at=None, break_end_at=None)
    assert invalid.status_code == 422


def test_annual_report_matches_months_and_groups_daily(client):
    test_client, _ = client
    employee, post = setup_people(client)
    for month in (1, 2):
        response = add_shift(test_client, employee, post, start_at=f"2026-{month:02d}-10T08:00", end_at=f"2026-{month:02d}-10T16:00", break_start_at=None, break_end_at=None)
        assert response.status_code == 201
    annual = test_client.get("/api/reports", params={"year": 2026, "group_by_employee": "false"}).json()
    january = test_client.get("/api/reports", params={"year": 2026, "month": 1}).json()
    february = test_client.get("/api/reports", params={"year": 2026, "month": 2}).json()
    assert annual["totals"]["minutes"] == january["totals"]["minutes"] + february["totals"]["minutes"]
    assert len(annual["days"]) == 2
    assert sum(day["shift_count"] for day in annual["days"]) == 2


def test_csv_import_autocreates_missing_master_data(client):
    test_client, _ = client
    employee, post = setup_people(client)
    csv_content = (
        "id,puesto,municipio,fecha,entrada,salida,pausa inicio,pausa fin,pausa remunerada\n"
        f"{employee['code']},{post['name']},{post['municipality']},2026-04-12,22:00,06:00,,,\n"
        "VS-NUEVO,Puesto Nuevo,València,2026-04-13,08:00,16:00,,,\n"
    ).encode()
    preview = test_client.post("/api/imports/csv/preview", files={"file": ("turnos.csv", csv_content, "text/csv")})
    assert preview.status_code == 200, preview.text
    assert preview.json()["ready"] == 2
    result = test_client.post("/api/imports/csv/confirm", json={"filename": "turnos.csv", "rows": preview.json()["rows"]})
    assert result.status_code == 201, result.text
    assert result.json()["pending"] == 2
    assert result.json()["needs_review"] == 0
    assert any(e["code"] == "VS-NUEVO" for e in test_client.get("/api/employees?limit=5000").json())
    assert any(p["name"] == "Puesto Nuevo" for p in test_client.get("/api/posts").json())
    report = test_client.get("/api/reports", params={"year": 2026, "month": 4}).json()
    assert report["totals"]["hours"] == "00:00"
    queue = test_client.get("/api/review").json()
    assert len(queue) == 2
    valid_pending = next(item for item in queue if item["kind"] == "shift")
    assert test_client.post(f"/api/shifts/{valid_pending['id']}/confirm").status_code == 200
    report = test_client.get("/api/reports", params={"year": 2026, "month": 4}).json()
    assert report["totals"]["hours"] == "08:00"


def test_csv_semicolon_separator_is_detected(client):
    test_client, _ = client
    employee, post = setup_people(client)
    csv_content = (
        "id;puesto;municipio;fecha;entrada;salida\n"
        f"{employee['code']};{post['name']};{post['municipality']};2026-04-12;22:00;06:00\n"
    ).encode()
    preview = test_client.post("/api/imports/preview", files={"file": ("turnos.csv", csv_content, "text/csv")})
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["status"] == "ready", row["errors"]
    assert row["start_at"] == "2026-04-12T22:00"
    assert row["end_at"] == "2026-04-13T06:00"


def test_exported_csv_can_be_reimported(client):
    test_client, _ = client
    employee, post = setup_people(client)
    assert add_shift(test_client, employee, post, start_at="2026-04-12T22:00", end_at="2026-04-13T06:00", break_start_at=None, break_end_at=None).status_code == 201
    export = test_client.get("/api/exports/shifts.csv", params={"year": 2026, "month": 4})
    assert export.status_code == 200
    assert export.text.splitlines()[0].lstrip("\ufeff").startswith("ID empleado;")
    preview = test_client.post("/api/imports/preview", files={"file": ("cuadrante-2026-04.csv", export.content, "text/csv")})
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["status"] == "ready", row["errors"]
    assert row["employee_id"] == employee["id"]
    assert row["post_id"] == post["id"]
    assert row["start_at"] == "2026-04-12T22:00"
    assert row["end_at"] == "2026-04-13T06:00"


def test_import_flags_same_name_different_id_for_review(client):
    test_client, _ = client
    employee, post = setup_people(client)
    csv_content = (
        "id,empleado,puesto,municipio,fecha,entrada,salida\n"
        f"VS-002,Ana Vigilante,{post['name']},{post['municipality']},2026-04-12,08:00,16:00\n"
    ).encode()
    preview = test_client.post("/api/imports/preview", files={"file": ("turnos.csv", csv_content, "text/csv")})
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["status"] == "pending"
    assert row["duplicate_of"] == ["VS-001"]


def test_duplicate_names_are_listed_and_mergeable(client):
    test_client, _ = client
    target = test_client.post("/api/employees", json={"code": "VS-001", "first_name": "Ana", "last_name": "Vigilante"}).json()
    source = test_client.post("/api/employees", json={"code": "VS-002", "first_name": "Ana", "last_name": "Vigilante"}).json()
    post = test_client.post("/api/posts", json={"name": "Puesto", "municipality": "Castelló de la Plana", "province": "Castellón"}).json()
    assert test_client.post("/api/shifts", json={"employee_id": source["id"], "post_id": post["id"], "start_at": "2026-04-12T08:00", "end_at": "2026-04-12T16:00", "break_paid": False}).status_code == 201
    duplicates = test_client.get("/api/duplicates").json()
    assert len(duplicates) == 1 and len(duplicates[0]["employees"]) == 2
    merged = test_client.post(f"/api/employees/{target['id']}/merge", json={"source_id": source["id"]}).json()
    assert merged["merged_shifts"] == 1
    assert test_client.get("/api/duplicates").json() == []
    report = test_client.get("/api/reports", params={"year": 2026, "month": 4}).json()
    assert report["totals"]["hours"] == "08:00"


def test_dst_ambiguous_and_nonexistent_local_times_rejected(client):
    test_client, _ = client
    employee, post = setup_people(client)
    nonexistent = add_shift(test_client, employee, post, start_at="2026-03-29T02:15", end_at="2026-03-29T03:15", break_start_at=None, break_end_at=None)
    ambiguous = add_shift(test_client, employee, post, start_at="2026-10-25T02:15", end_at="2026-10-25T03:15", break_start_at=None, break_end_at=None)
    assert nonexistent.status_code == 422
    assert ambiguous.status_code == 422


def test_fast_segment_calculator_and_exact_boundary_rules():
    shift = {
        "start_at": "2026-04-10T21:30", "end_at": "2026-04-11T06:30",
        "break_start_at": "2026-04-11T01:00", "break_end_at": "2026-04-11T01:30", "break_paid": 0,
    }
    metrics = shift_metrics(shift, {"2026-04-11"})
    assert metrics["minutes"] == 510
    assert metrics["night_minutes"] == 450
    assert metrics["holiday_minutes"] == 360
    assert metrics["days"]["2026-04-10"]["minutes"] == 150
    assert metrics["days"]["2026-04-11"]["minutes"] == 360


def test_xlsx_import_autocreates_missing_rows(client):
    import openpyxl

    test_client, _ = client
    employee, post = setup_people(client)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["id", "puesto", "municipio", "fecha", "entrada", "salida", "pausa inicio", "pausa fin", "pausa remunerada"])
    sheet.append([employee["code"], post["name"], post["municipality"], "2026-04-12", "22:00", "06:00", None, None, None])
    sheet.append(["VS-NUEVO", "Puesto Nuevo", "València", "2026-04-13", "08:00", "16:00", None, None, None])
    buffer = io.BytesIO()
    workbook.save(buffer)

    preview = test_client.post("/api/imports/preview", files={"file": ("turnos.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert preview.status_code == 200, preview.text
    assert preview.json()["kind"] == "xlsx"
    result = test_client.post("/api/imports/confirm", json={"filename": "turnos.xlsx", "source": "xlsx", "rows": preview.json()["rows"]})
    assert result.status_code == 201, result.text
    assert result.json()["pending"] == 2
    assert result.json()["needs_review"] == 0
    report = test_client.get("/api/reports", params={"year": 2026, "month": 4}).json()
    assert report["totals"]["hours"] == "00:00"


def test_official_holidays_are_loaded_for_default_municipality(client):
    test_client, _ = client
    rows_2026 = test_client.get("/api/holidays", params={"year": 2026, "municipality": "Castelló de la Plana"}).json()
    dates = {row["date"] for row in rows_2026}
    assert "2026-03-19" in dates  # San José (autonómico)
    assert "2026-03-09" in dates  # Fiestas de la Magdalena (local Castelló)
    assert "2026-06-29" in dates  # San Pedro (local Castelló)
    assert "2026-10-09" in dates  # Día de la Comunitat Valenciana
    assert "2026-12-25" in dates  # Navidad
    second = test_client.get("/api/holidays", params={"year": 2026, "municipality": "Castelló de la Plana"}).json()
    assert len(second) == len(rows_2026)  # carga idempotente


def test_monthly_summary_is_stored_apart_from_shift_totals(client):
    test_client, _ = client
    csv_content = (
        "ID;Nombre ficticio;Puesto;Localización ficticia;Mes;Horas totales trabajadas\n"
        "SEG-001;Álvaro Castro;Vigilante de seguridad;Centro Comercial Aurora;July 2026;200\n"
    ).encode()
    preview = test_client.post("/api/imports/preview", files={"file": ("resumen.csv", csv_content, "text/csv")})
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["kind"] == "summary", row["errors"]
    assert row["month"] == "2026-07"
    assert row["minutes"] == 200 * 60
    result = test_client.post("/api/imports/confirm", json={"filename": "resumen.csv", "rows": preview.json()["rows"]})
    assert result.status_code == 201, result.text
    assert result.json()["summaries"] == 1
    assert result.json()["pending"] == 0
    report = test_client.get("/api/reports", params={"year": 2026, "month": 7}).json()
    assert report["totals"]["hours"] == "00:00"  # no es un turno
    assert report["summaries"][0]["hours"] == "200:00"
    summaries = test_client.get("/api/monthly-summaries", params={"year": 2026}).json()
    assert len(summaries) == 1 and summaries[0]["employee_code"] == "SEG-001"


def test_xlsx_datetime_cells_are_interpreted(client):
    import openpyxl

    test_client, _ = client
    employee, post = setup_people(client)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["id", "puesto", "municipio", "fecha", "entrada", "salida"])
    sheet.append([employee["code"], post["name"], post["municipality"], datetime(2026, 4, 12, 0, 0), datetime(2026, 4, 12, 22, 0), datetime(2026, 4, 13, 6, 0)])
    buffer = io.BytesIO()
    workbook.save(buffer)

    preview = test_client.post("/api/imports/preview", files={"file": ("turnos.xlsx", buffer.getvalue(), "application/vnd.ms-excel")})
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["status"] == "ready", row["errors"]
    assert row["start_at"] == "2026-04-12T22:00"
    assert row["end_at"] == "2026-04-13T06:00"


def _pdf_with_table(headers, rows):
    import fitz

    document = fitz.open()
    page = document.new_page(width=595, height=200 + 22 * len(rows))
    x0, y0, colw, rowh = 40, 40, 90, 22
    for r, row in enumerate([headers, *rows]):
        for c, cell in enumerate(row):
            rect = fitz.Rect(x0 + c * colw, y0 + r * rowh, x0 + (c + 1) * colw, y0 + (r + 1) * rowh)
            page.draw_rect(rect, color=(0, 0, 0), width=0.6)
            page.insert_text((rect.x0 + 3, rect.y0 + 15), str(cell), fontsize=9)
    data = document.tobytes()
    document.close()
    return data


def test_pdf_table_import(client):
    test_client, _ = client
    employee, post = setup_people(client)
    data = _pdf_with_table(
        ["id", "puesto", "municipio", "fecha", "entrada", "salida"],
        [[employee["code"], post["name"], post["municipality"], "2026-04-12", "22:00", "06:00"]],
    )
    preview = test_client.post("/api/imports/preview", files={"file": ("cuadrante.pdf", data, "application/pdf")})
    assert preview.status_code == 200, preview.text
    assert preview.json()["kind"] == "pdf"
    row = preview.json()["rows"][0]
    assert row["status"] == "ready", row["errors"]
    assert row["start_at"] == "2026-04-12T22:00"


def test_pdf_without_table_is_rejected(client):
    import fitz

    test_client, _ = client
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Documento sin tabla")
    data = document.tobytes()
    document.close()
    response = test_client.post("/api/imports/preview", files={"file": ("escaneado.pdf", data, "application/pdf")})
    assert response.status_code == 422


def test_import_rejects_unsupported_extension_and_broken_signature(client):
    test_client, _ = client
    docx = test_client.post("/api/imports/preview", files={"file": ("datos.docx", b"x", "application/octet-stream")})
    assert docx.status_code == 415
    fake_xlsx = test_client.post("/api/imports/preview", files={"file": ("falso.xlsx", b"no es un zip", "application/vnd.ms-excel")})
    assert fake_xlsx.status_code == 422
    fake_pdf = test_client.post("/api/imports/preview", files={"file": ("falso.pdf", b"no es pdf", "application/pdf")})
    assert fake_pdf.status_code == 422


def test_api_responses_are_not_cached(client):
    test_client, _ = client
    response = test_client.get("/api/employees")
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"


def test_confirmed_shifts_can_be_sent_to_review_individually_and_in_bulk(client):
    test_client, _ = client
    employee, post = setup_people(client)
    first = add_shift(test_client, employee, post)
    second = add_shift(test_client, employee, post, start_at="2026-04-11T08:00", end_at="2026-04-11T16:00", break_start_at=None, break_end_at=None)
    assert first.status_code == 201 and second.status_code == 201

    response = test_client.post("/api/shifts/review", json={"ids": [first.json()["id"], second.json()["id"], "no-existe"]})
    assert response.status_code == 200
    assert response.json() == {"updated": 2, "requested": 3}

    report = test_client.get("/api/reports", params={"year": 2026, "month": 4}).json()
    assert report["totals"]["hours"] == "00:00"
    assert report["shifts"] == []

    queue = test_client.get("/api/review").json()
    pending_ids = {item["shift"]["id"] for item in queue if item["kind"] == "shift"}
    assert pending_ids == {first.json()["id"], second.json()["id"]}

    again = test_client.post("/api/shifts/review", json={"ids": [first.json()["id"]]})
    assert again.json()["updated"] == 0


def test_future_holidays_are_generated_from_rules(client):
    test_client, _ = client
    rows_2030 = test_client.get("/api/holidays", params={"year": 2030, "municipality": "Castelló de la Plana"}).json()
    dates = {row["date"] for row in rows_2030}
    assert "2030-01-01" in dates  # Año Nuevo
    assert "2030-03-19" in dates  # San José (autonómico)
    assert "2030-04-19" in dates  # Viernes Santo
    assert "2030-04-22" in dates  # Lunes de Pascua
    assert "2030-03-25" in dates  # Magdalena (local calculada)
    assert "2030-06-29" in dates  # San Pedro (local)
    assert "2030-10-09" in dates  # Comunitat Valenciana
    assert "2030-12-25" in dates  # Navidad

    dates_2031 = {row["date"] for row in test_client.get("/api/holidays", params={"year": 2031}).json()}
    assert "2031-06-30" in dates_2031  # San Pedro cae domingo: se traslada al lunes
    assert "2031-06-29" not in dates_2031


def test_exported_xlsx_can_be_reimported(client):
    import openpyxl

    test_client, _ = client
    employee, post = setup_people(client)
    assert add_shift(test_client, employee, post, start_at="2026-04-12T22:00", end_at="2026-04-13T06:00", break_start_at=None, break_end_at=None).status_code == 201

    export = test_client.get("/api/exports/shifts.xlsx", params={"year": 2026, "month": 4})
    assert export.status_code == 200
    assert export.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    workbook = openpyxl.load_workbook(io.BytesIO(export.content), read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    assert rows[0][0] == "ID empleado"
    assert rows[1][5] == "2026-04-12"
    assert rows[1][6] == "22:00"
    assert rows[1][7] == "06:00"

    preview = test_client.post("/api/imports/preview", files={"file": ("cuadrante-2026-04.xlsx", export.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["status"] == "ready", row["errors"]
    assert row["employee_id"] == employee["id"]
    assert row["post_id"] == post["id"]
    assert row["start_at"] == "2026-04-12T22:00"
    assert row["end_at"] == "2026-04-13T06:00"


def test_csv_windows_encoding_title_row_and_spanish_date_are_accepted(client):
    test_client, _ = client
    employee, post = setup_people(client)
    content = (
        "LISTADO DE TURNOS\n\n"
        "ID;Empleado;Apellidos;Puesto;Municipio;Fecha;Hora entrada;Hora salida\n"
        f"{employee['code']};Álvaro;Vigilante;{post['name']};{post['municipality']};12/04/2026;22:00;06:00\n"
    ).encode("cp1252")
    preview = test_client.post("/api/imports/preview", files={"file": ("turnos.csv", content, "text/csv")})
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["status"] == "ready", row["errors"]
    assert row["employee_name"] == "Álvaro Vigilante"
    assert row["start_at"] == "2026-04-12T22:00"
    assert row["end_at"] == "2026-04-13T06:00"


def test_xlsx_finds_header_after_title_and_reads_excel_serial_date(client):
    import openpyxl

    test_client, _ = client
    employee, post = setup_people(client)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["CUADRANTE MENSUAL"])
    sheet.append([])
    sheet.append(["ID", "Nombre", "Apellidos", "Puesto", "Municipio", "Fecha", "Entrada", "Salida"])
    sheet.append([employee["code"], "Ana", "Vigilante", post["name"], post["municipality"], 46124, 0.5, 0.75])
    buffer = io.BytesIO()
    workbook.save(buffer)

    preview = test_client.post("/api/imports/preview", files={"file": ("turnos.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["status"] == "ready", row["errors"]
    assert row["employee_name"] == "Ana Vigilante"
    assert row["start_at"] == "2026-04-12T12:00"
    assert row["end_at"] == "2026-04-12T18:00"


def test_national_holiday_applies_to_non_default_municipality(client):
    test_client, _ = client
    employee, post = setup_people(client)  # puesto en València
    shift = add_shift(test_client, employee, post, start_at="2026-01-01T08:00", end_at="2026-01-01T16:00", break_start_at=None, break_end_at=None)
    assert shift.status_code == 201
    report = test_client.get("/api/reports", params={"year": 2026, "month": 1}).json()
    assert report["totals"]["holiday_hours"] == "08:00"
    valencia = test_client.get("/api/holidays", params={"year": 2026, "municipality": "València"}).json()
    dates = {row["date"] for row in valencia}
    assert "2026-01-01" in dates
    assert "2026-03-09" not in dates  # Magdalena es local de Castelló


def test_import_end_datetime_is_honored_and_32h_is_rejected(client):
    test_client, _ = client
    employee, post = setup_people(client)
    good = (
        "id;puesto;municipio;entrada;salida\n"
        f"{employee['code']};{post['name']};{post['municipality']};2026-04-12T22:00;2026-04-13T06:00\n"
    ).encode()
    preview = test_client.post("/api/imports/preview", files={"file": ("turnos.csv", good, "text/csv")})
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["status"] == "ready", row["errors"]
    assert row["end_at"] == "2026-04-13T06:00"

    long = (
        "id;puesto;municipio;entrada;salida\n"
        f"{employee['code']};{post['name']};{post['municipality']};2026-04-12T08:00;2026-04-13T16:00\n"
    ).encode()
    preview = test_client.post("/api/imports/preview", files={"file": ("largo.csv", long, "text/csv")})
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["status"] == "pending"
    assert any("24 horas" in error for error in row["errors"])


def test_summary_invalid_hours_do_not_crash(client):
    test_client, _ = client
    employee, _ = setup_people(client)
    content = (
        "id;mes;horas totales\n"
        f"{employee['code']};2026-07;-5\n"
        f"{employee['code']};2026-08;inf\n"
    ).encode()
    preview = test_client.post("/api/imports/preview", files={"file": ("resumen.csv", content, "text/csv")})
    assert preview.status_code == 200, preview.text
    assert preview.json()["pending"] == 2
    assert all(row["status"] == "pending" for row in preview.json()["rows"])


def test_xlsx_multiple_sheets_are_combined(client):
    import openpyxl

    test_client, _ = client
    employee, post = setup_people(client)
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "Enero"
    first.append(["id", "puesto", "municipio", "fecha", "entrada", "salida"])
    first.append([employee["code"], post["name"], post["municipality"], "2026-04-12", "08:00", "16:00"])
    second = workbook.create_sheet("Febrero")
    second.append(["id", "puesto", "municipio", "fecha", "entrada", "salida"])
    second.append([employee["code"], post["name"], post["municipality"], "2026-04-13", "08:00", "16:00"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    preview = test_client.post("/api/imports/preview", files={"file": ("dos-hojas.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert preview.status_code == 200, preview.text
    assert preview.json()["count"] == 2


def test_deleted_generated_holiday_survives_restart(tmp_path):
    db_path = tmp_path / "reinicio.sqlite3"
    app = create_app(db_path)
    with TestClient(app) as test_client:
        rows = test_client.get("/api/holidays", params={"year": 2030, "municipality": "Castelló de la Plana"}).json()
        holiday = next(row for row in rows if row["date"] == "2030-01-01")
        assert test_client.delete(f"/api/holidays/{holiday['id']}").status_code == 204
        rows = test_client.get("/api/holidays", params={"year": 2030, "municipality": "Castelló de la Plana"}).json()
        assert all(row["date"] != "2030-01-01" for row in rows)
    with TestClient(create_app(db_path)) as restarted:
        rows = restarted.get("/api/holidays", params={"year": 2030, "municipality": "Castelló de la Plana"}).json()
        assert all(row["date"] != "2030-01-01" for row in rows)


def test_summary_review_can_be_corrected_without_creating_a_shift(client):
    test_client, _ = client
    employee, _ = setup_people(client)
    content = (
        "id;mes;horas totales\n"
        f"{employee['code']};2026-13;200\n"
    ).encode()
    preview = test_client.post("/api/imports/preview", files={"file": ("resumen.csv", content, "text/csv")})
    assert preview.status_code == 200, preview.text
    result = test_client.post("/api/imports/confirm", json={"filename": "resumen.csv", "rows": preview.json()["rows"]})
    assert result.status_code == 201, result.text
    queue = test_client.get("/api/review").json()
    item = next(entry for entry in queue if entry.get("raw", {}).get("kind") == "summary")
    resolved = test_client.post(f"/api/review/{item['id']}/resolve-summary", json={"employee_id": employee["id"], "month": "2026-07", "hours": "200"})
    assert resolved.status_code == 201, resolved.text
    summaries = test_client.get("/api/monthly-summaries", params={"year": 2026, "month": 7}).json()
    assert len(summaries) == 1
    assert summaries[0]["minutes"] == 12000
    assert test_client.get("/api/reports", params={"year": 2026, "month": 4}).json()["shifts"] == []


def test_excel_summary_accepts_timedelta_and_datetime_month(client):
    import openpyxl
    from datetime import datetime, timedelta

    test_client, _ = client
    employee, _ = setup_people(client)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["id", "mes", "horas totales"])
    sheet.append([employee["code"], datetime(2026, 7, 1), timedelta(hours=200)])
    buffer = io.BytesIO()
    workbook.save(buffer)
    preview = test_client.post("/api/imports/preview", files={"file": ("resumen.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["status"] == "ready", row["errors"]
    assert row["month"] == "2026-07"
    assert row["minutes"] == 12000


def test_duplicate_new_names_in_same_file_are_flagged(client):
    test_client, _ = client
    _, post = setup_people(client)
    content = (
        "id;empleado;puesto;municipio;fecha;entrada;salida\n"
        f"VS-N1;Ana Nueva;{post['name']};{post['municipality']};2026-04-12;08:00;16:00\n"
        f"VS-N2;Ana Nueva;{post['name']};{post['municipality']};2026-04-13;08:00;16:00\n"
    ).encode()
    preview = test_client.post("/api/imports/preview", files={"file": ("duplicados.csv", content, "text/csv")})
    assert preview.status_code == 200, preview.text
    rows = preview.json()["rows"]
    assert rows[0]["status"] == "ready"
    assert rows[1]["status"] == "pending"
    assert rows[1]["duplicate_of"] == ["VS-N1"]


def test_export_neutralizes_formula_injection(client):
    import openpyxl

    test_client, _ = client
    employee = test_client.post("/api/employees", json={"code": "=CMD", "first_name": "=1+1", "last_name": "Peligro"}).json()
    post = test_client.post("/api/posts", json={"name": "=Puesto", "municipality": "València", "province": "Valencia"}).json()
    assert add_shift(test_client, employee, post).status_code == 201

    csv_export = test_client.get("/api/exports/shifts.csv", params={"year": 2026, "month": 4})
    assert "'=CMD" in csv_export.text
    xlsx_export = test_client.get("/api/exports/shifts.xlsx", params={"year": 2026, "month": 4})
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx_export.content), read_only=True)
    row = list(workbook.active.iter_rows(values_only=True))[1]
    workbook.close()
    assert row[0].startswith("'=")
    assert row[1].startswith("'=")
    assert row[2].startswith("'=")
